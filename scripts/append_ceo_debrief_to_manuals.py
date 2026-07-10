import docx
import os

def append_debrief_to_doc(filepath, is_facilitator=False):
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        return

    doc = docx.Document(filepath)
    doc.add_page_break()
    
    if is_facilitator:
        # --- Section 18: 2026 Platform Updates ---
        doc.add_heading('Section 18: Summary of 2026 Platform Updates', level=1)
        doc.add_paragraph(
            "During the 2026 systems re-audit, several core engine enhancements, bug fixes, and "
            "facilitator orchestration tools were deployed to improve classroom reliability and pedagogical engagement."
        )
        
        doc.add_heading('18.1 Teachable Moments Detector (B5)', level=2)
        doc.add_paragraph(
            "The Teachable Moments engine is an automated, read-only facilitator alert system. "
            "It continuously monitors active decision flags across the cohort. If 50% or more of the active teams "
            "converge on an adverse decision flag (such as the Round 1 'electronics_blindspot' flag or "
            "the Round 2 'materiality_ignored' flag), a private alert panel appears on the Facilitator Dashboard. "
            "This panel indicates the source round, target round, and specific impact of the flag, warning "
            "the facilitator to pause the cohort and debrief the structural risks before the consequences manifest."
        )
        
        doc.add_heading('18.2 Round 2 Stakeholder Panel Survey (Multi-Group Model)', level=2)
        doc.add_paragraph(
            "Round 2 has transitioned to a Multi-Group Model for stakeholder survey commission. "
            "Instead of a single tiered survey fee, facilitators and players can now commission "
            "independent stakeholder panel surveys at a flat fee of $750,000 USD per group. "
            "Four stakeholder groups are available: (1) Investors, focusing on financial and transition risks; "
            "(2) Own Workforce, focusing on internal working conditions; (3) NGOs & Communities, focusing on "
            "environmental harms; and (4) Subject Matter Experts, validating ESRS topic mapping and materiality thresholds. "
            "The legacy single tiered survey is still supported for backward compatibility."
        )
        
        doc.add_heading('18.3 Single-BU Vertical Alignment', level=2)
        doc.add_paragraph(
            "A bug in the session creation engine was resolved where sessions configured in Single Business Unit mode "
            "with a substituted industry vertical (e.g. retail_fmcg, oil_gas) incorrectly granted players the full "
            "4-BU conglomerate. The engine now correctly limits the cohort to exactly 1 vertical BU carrying the "
            "selected sector profile."
        )
        
        doc.add_heading('18.4 Password Reset Policy', level=2)
        doc.add_paragraph(
            "To support data privacy and cohort security, players are flagged with a forced password change "
            "on first login. Once they set a personal password in the cockpit, the 'must_change_password' flag "
            "is automatically cleared in the session database and metadata."
        )

        doc.add_page_break()

        # --- CEO Debrief Section ---
        doc.add_heading('Section: The CEO Debrief & Assessment Engine', level=1)
        doc.add_paragraph(
            "The Muressons Global Command simulation concludes with a high-fidelity, pedagogically rigorous CEO Debrief module. "
            "This engine moves beyond deterministic point-in-time scoring by leveraging a live LLM integration and multi-dimensional analysis "
            "to assess leadership competency."
        )
        
        doc.add_heading('1. Blended Qualitative & Quantitative Scoring', level=2)
        doc.add_paragraph(
            "Every player is evaluated across six core dimensions (e.g., Strategic Thinking, Adaptive Leadership). "
            "The final score is a 50/50 blend of Data Performance (derived objectively from their 10-round choices) "
            "and Interview Performance (qualitative assessment of their verbal/textual responses by an AI Persona, "
            "such as GPT-4o or Claude 3.5)."
        )
        
        doc.add_heading('2. Trajectory-Aware Analysis', level=2)
        doc.add_paragraph(
            "The scoring engine does not just look at the final Round 10 results. It analyzes the entire time-series trajectory "
            "of the player's decisions. Consistent improvement or stability yields a positive modifier (up to +15%), "
            "while high volatility or late-game collapses trigger penalties. This rewards sustainable, deliberate strategic behavior."
        )
        
        doc.add_heading('3. Evidence Citations', level=2)
        doc.add_paragraph(
            "To prevent generic feedback, the assessment engine mines the historical decision log. When presenting the debrief, "
            "it cites specific rounds and actual player choices (e.g., 'In Round 5, you chose to increase R&D by $8M...'). "
            "This grounds the debrief in objective reality."
        )
        
        doc.add_heading('4. Adaptive Questioning', level=2)
        doc.add_paragraph(
            "The system identifies the two lowest-scoring dimensions from the player's data score and dynamically generates "
            "probe questions targeting those weak spots. This forces the learner to reflect precisely where they struggled the most."
        )
        
        doc.add_heading('5. Metacognitive Self-Assessment & Calibration', level=2)
        doc.add_paragraph(
            "Before the debrief interview begins, players rate their own performance from 1-10 on all six dimensions. "
            "The engine compares this self-assessment to their actual performance score to detect calibration gaps, "
            "flagging blind spots such as 'Significantly Overconfident' or 'Underconfident'."
        )
        
        doc.add_heading('6. Counterfactual What-If Exploration', level=2)
        doc.add_paragraph(
            "Facilitators can trigger 'What-If' analyses, allowing the engine to calculate hypothetical score deltas if the player "
            "had made a specific decision earlier in the simulation, reinforcing the compounding nature of early interventions."
        )

    else:
        # Student Version
        doc.add_heading('Appendix: The CEO Debrief', level=1)
        doc.add_paragraph(
            "At the conclusion of Round 10, your journey does not end. You will face a formal CEO Debrief—a comprehensive "
            "competency assessment designed to evaluate your leadership capability beyond raw financial metrics."
        )
        
        doc.add_heading('How You Are Evaluated', level=2)
        doc.add_paragraph("Your final competency score across six core dimensions is calculated through a blended approach:", style='List Bullet')
        doc.add_paragraph("50% Performance Data: Calculated directly from your decisions, outcomes, and long-term consistency (trajectory) across all 10 rounds.", style='List Bullet')
        doc.add_paragraph("50% Interview Responses: You will answer a series of adaptive questions during the debrief. Your ability to articulate trade-offs, rationale, and systemic awareness is evaluated by an AI CEO persona.", style='List Bullet')
        
        doc.add_heading('Metacognitive Self-Assessment', level=2)
        doc.add_paragraph(
            "Before the interview begins, you will be asked to rate your own performance on a scale of 1 to 10 for each dimension. "
            "Be honest. The simulation will measure your 'Calibration Gap'—the difference between how you perceive your performance "
            "and how the data actually reflects it. High self-awareness is a critical executive trait."
        )
        
        doc.add_heading('Evidence-Based Feedback', level=2)
        doc.add_paragraph(
            "The debrief is highly personalized. Expect the CEO to pull 'receipts'—citing specific decisions you made in specific rounds "
            "to justify the feedback provided. You will also see how your final scores compare to your cohort via Peer Benchmarking."
        )

    # Save logic
    dir_name = os.path.dirname(filepath)
    base_name = os.path.basename(filepath)
    name, ext = os.path.splitext(base_name)
    
    # Simple versioning for naming
    if "with_CEO_Debrief" not in name:
        new_filepath = os.path.join(dir_name, f"{name}_with_CEO_Debrief{ext}")
    else:
        new_filepath = filepath
        
    doc.save(new_filepath)
    print(f"Successfully appended updates and CEO Debrief to {new_filepath}")

if __name__ == '__main__':
    base_dir = r"c:\Users\Home\.gemini\antigravity\scratch\muressons-sim"
    
    # Target latest manuals
    facilitator_manual = os.path.join(base_dir, "Muressons_Simulation_Briefings ver 11.docx")
    append_debrief_to_doc(facilitator_manual, is_facilitator=True)
