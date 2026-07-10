import os
import docx
from docx.shared import Pt, Cm

BASE_DIR = r"c:\Users\Home\.gemini\antigravity\scratch\muressons-sim"

FACILITATOR_MANUAL = os.path.join(BASE_DIR, "Muressons_Facilitator_Manual_v5.docx")
STUDENT_MANUAL = os.path.join(BASE_DIR, "Muressons_Student_Manual_v5_with_CAROIC_with_CEO_Debrief_with_Constellation.docx")

def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)

def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.style.font.size = Pt(11)
    return p

def add_bullet(doc, text):
    doc.add_paragraph(text, style='List Bullet')

def add_materiality_calculations_appendix(doc):
    doc.add_page_break()

    # Title
    add_heading(doc, 'Appendix: Double Materiality Matrix Calculations', level=1)
    add_body(doc,
        "This appendix details the calculations behind the Double Materiality Matrix in Round 2, "
        "explaining how the unlocked budget, stakeholder panel fees, and mitigation costs are derived."
    )

    # 1. Unlocked Budget
    add_heading(doc, '1. Unlocked Budget', level=2)
    add_body(doc,
        "When submitting the materiality matrix, the system calculates the capital unlocked based on "
        "accuracy in Quadrant 1 (High Financial & High Societal Impact)."
    )
    add_bullet(doc, "Total Available Budget: The base materiality budget defaults to $15,000,000.")
    add_bullet(doc, "Accuracy Calculation (m_acc): The system evaluates the number of true Q1 issues correctly placed in Q1, divided by the total number of true Q1 issues.")
    add_bullet(doc, "Formula: Allocated Budget = Total Available Budget × Accuracy")
    add_bullet(doc, "Clawback Penalty: If Option C (Ignore Framework) was selected in the governance decision, the backend applies a 40% budget clawback to the unlocked amount to reflect the misalignment between the governance posture and the materiality exercise.")

    # 2. Stakeholder Panel / Consultant Fee
    add_heading(doc, '2. Stakeholder Panel / Consultant Fee', level=2)
    add_body(doc,
        "The cost to 'Commission Stakeholder Panel Survey' scales dynamically based on how many issues are pre-rated."
    )
    add_bullet(doc, "Issues 1-4: Base fee of $250,000 per issue.")
    add_bullet(doc, "Issues 5-8: Extended fee of $500,000 per issue (the cost doubles beyond 4 issues to reflect the complexity of wider engagement).")
    add_body(doc, "Examples:")
    add_bullet(doc, "1 Issue = $250,000")
    add_bullet(doc, "4 Issues = $1,000,000 ($250K × 4)")
    add_bullet(doc, "8 Issues = $3,000,000 ($1M for first four + $2M for the next four)")

    # 3. Mitigation Costs
    add_heading(doc, '3. Mitigation Costs (Issue Chips)', level=2)
    add_body(doc,
        "When an issue chip is dragged on the matrix, it displays a specific cost (e.g., '$1.2M')."
    )
    add_bullet(doc, "Source: This number is stored as 'mitigation_cost_usd' within each issue's configuration in the backend.")
    add_bullet(doc, "Prioritise Budget: The sum of these costs for all issues placed in the Q1 container forms the 'Prioritise Budget' shown dynamically at the top of the matrix. The numbers are formatted using division by 1,000,000 to display them with the 'M' suffix.")

def process(filepath, add_fn, label, output_name):
    if not os.path.exists(filepath):
        print(f"[ERROR] File not found: {filepath}")
        return

    doc = docx.Document(filepath)
    add_fn(doc)

    out = os.path.join(BASE_DIR, output_name)
    doc.save(out)
    print(f"[OK] {label} saved -> {output_name}")

if __name__ == '__main__':
    process(
        FACILITATOR_MANUAL,
        add_materiality_calculations_appendix,
        "Facilitator Manual",
        "Muressons_Facilitator_Manual_v6.docx",
    )
    process(
        STUDENT_MANUAL,
        add_materiality_calculations_appendix,
        "Student Manual",
        "Muressons_Student_Manual_v6.docx",
    )
