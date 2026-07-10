"""
Student Manual - Shared helpers for Word document generation.
"""
import re, html
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

NAVY = RGBColor(0x0f, 0x17, 0x2a)
TEAL = RGBColor(0x0d, 0x9b, 0x8c)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
LIGHT_BG = RGBColor(0xf0, 0xf4, 0xf8)
RED_ACCENT = RGBColor(0xc0, 0x39, 0x2b)
GOLD = RGBColor(0xd4, 0x9a, 0x2a)

def strip_html(text):
    text = re.sub(r'<[^>]+>', '', str(text))
    return html.unescape(text).replace("\\n", "\n")

def create_doc():
    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    return doc

def h(doc, text, level=1, color=NAVY):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = color
    return heading

def body(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(strip_html(text))
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = DARK_GRAY
    return p

def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(11)
        r = p.add_run(strip_html(text))
        r.font.size = Pt(11)
    else:
        r = p.add_run(strip_html(text))
        r.font.size = Pt(11)
    return p

def add_image(doc, path, width=Inches(5.8), caption=None):
    try:
        doc.add_picture(path, width=width)
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(caption)
            r.font.size = Pt(9)
            r.italic = True
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    except Exception as e:
        body(doc, f"[Screenshot: {caption or path} — {e}]", italic=True)

def shade_cells(row, color_hex="F0F4F8"):
    for cell in row.cells:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

def styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    for i, label in enumerate(headers):
        hdr[i].text = label
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    for row_data in rows:
        row = table.add_row().cells
        for i, val in enumerate(row_data):
            row[i].text = str(val)
            for p in row[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return table

def option_table(doc, options):
    headers = ['Option', 'Title', 'Description', 'Key Impacts']
    rows = []
    for key in sorted(options.keys()):
        opt = options[key]
        impacts = opt.get('impacts', {})
        impact_lines = []
        for k, v in impacts.items():
            if isinstance(v, bool):
                impact_lines.append(f"{k}: {'Yes' if v else 'No'}")
            elif isinstance(v, (int, float)):
                prefix = '+' if v > 0 else ''
                if abs(v) >= 1_000_000:
                    impact_lines.append(f"{k}: {prefix}${v/1_000_000:.0f}M")
                else:
                    impact_lines.append(f"{k}: {prefix}{v}")
        rows.append((
            opt.get('label', key[-1].upper()),
            opt.get('title', ''),
            strip_html(opt.get('description', '')),
            '\n'.join(impact_lines) if impact_lines else 'See description'
        ))
    return styled_table(doc, headers, rows)

def formula_block(doc, name, formula, desc, examples):
    h(doc, name, level=3)
    p = doc.add_paragraph()
    r = p.add_run('Formula: ')
    r.bold = True
    r.font.size = Pt(11)
    r = p.add_run(formula)
    r.font.name = 'Consolas'
    r.font.size = Pt(10)
    body(doc, desc)
    body(doc, 'Worked Example:', bold=True)
    for line in examples:
        bullet(doc, line)
    doc.add_paragraph()

def spacer(doc):
    doc.add_paragraph()

def page_break(doc):
    doc.add_page_break()
