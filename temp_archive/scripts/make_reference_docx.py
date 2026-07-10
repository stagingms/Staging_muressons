"""
Generate scripts/reference.docx — Muressons brand style template for pandoc.
Pandoc uses this file's styles when converting Markdown → DOCX.
Run: python scripts/make_reference_docx.py
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(__file__), "reference.docx")


def set_font(style, name, size_pt, bold=False, color=None):
    style.font.name = name
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    if color:
        style.font.color.rgb = color


def add_style_if_missing(doc, name, base_style_name="Normal"):
    """Return existing style or create a new paragraph style."""
    if name in [s.name for s in doc.styles]:
        return doc.styles[name]
    style = doc.styles.add_style(name, 1)  # 1 = WD_STYLE_TYPE.PARAGRAPH
    style.base_style = doc.styles[base_style_name]
    return style


def build_reference_doc():
    doc = Document()

    # ── Page margins ────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)

    # ── Normal / body text ───────────────────────────────────────────────────
    normal = doc.styles["Normal"]
    set_font(normal, "Calibri", 11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = Pt(14)

    # ── Heading 1 — Deep navy ────────────────────────────────────────────────
    h1 = doc.styles["Heading 1"]
    set_font(h1, "Calibri", 18, bold=True, color=RGBColor(0, 43, 91))
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after  = Pt(6)
    h1.paragraph_format.keep_with_next = True

    # ── Heading 2 — Corporate blue ───────────────────────────────────────────
    h2 = doc.styles["Heading 2"]
    set_font(h2, "Calibri", 14, bold=True, color=RGBColor(0, 90, 156))
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after  = Pt(4)
    h2.paragraph_format.keep_with_next = True

    # ── Heading 3 — Teal ────────────────────────────────────────────────────
    h3 = doc.styles["Heading 3"]
    set_font(h3, "Calibri", 12, bold=True, color=RGBColor(31, 117, 140))
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after  = Pt(4)
    h3.paragraph_format.keep_with_next = True

    # ── Heading 4 ────────────────────────────────────────────────────────────
    h4 = doc.styles["Heading 4"]
    set_font(h4, "Calibri", 11, bold=True, color=RGBColor(60, 60, 60))
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after  = Pt(2)

    # ── Block Text (blockquote) ───────────────────────────────────────────────
    block = add_style_if_missing(doc, "Block Text")
    set_font(block, "Calibri", 11, color=RGBColor(80, 80, 80))
    block.font.italic = True
    block.paragraph_format.left_indent  = Inches(0.5)
    block.paragraph_format.space_before = Pt(4)
    block.paragraph_format.space_after  = Pt(4)

    # ── Source Code (fenced code blocks) ────────────────────────────────────
    code = add_style_if_missing(doc, "Source Code")
    set_font(code, "Courier New", 9)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after  = Pt(2)
    # Gray background via paragraph shading
    pPr = code.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "F4F4F4")
    pPr.append(shd)

    # ── Verbatim Char (inline code) ──────────────────────────────────────────
    if "Verbatim Char" not in [s.name for s in doc.styles]:
        vc = doc.styles.add_style("Verbatim Char", 2)  # 2 = WD_STYLE_TYPE.CHARACTER
        vc.font.name = "Courier New"
        vc.font.size = Pt(10)

    # ── List styles ───────────────────────────────────────────────────────────
    lb = doc.styles["List Bullet"]
    set_font(lb, "Calibri", 11)
    lb.paragraph_format.space_after = Pt(3)

    ln = doc.styles["List Number"]
    set_font(ln, "Calibri", 11)
    ln.paragraph_format.space_after = Pt(3)

    # ── Table Grid style ─────────────────────────────────────────────────────
    # Pandoc uses "Table" style; set compact spacing
    if "Table" in [s.name for s in doc.styles]:
        tbl_style = doc.styles["Table"]
        set_font(tbl_style, "Calibri", 10)
        tbl_style.paragraph_format.space_before = Pt(2)
        tbl_style.paragraph_format.space_after  = Pt(2)

    # ── Add a small watermark paragraph so the file is non-empty ────────────
    p = doc.add_paragraph()
    run = p.add_run("Muressons Reference Document — do not edit directly.")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(200, 200, 200)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print(f"reference.docx written to: {OUT}")


if __name__ == "__main__":
    build_reference_doc()
