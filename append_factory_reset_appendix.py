"""
append_factory_reset_appendix.py
---------------------------------
Appends a detailed Factory Reset & Data Deletion Reference appendix to:
  - Muressons_Facilitator_Manual_v4_with_Constellation.docx  (Facilitator Manual)

Outputs versioned copy:
  - Muressons_Facilitator_Manual_v5.docx
"""

import os
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE_DIR = r"c:\Users\Home\.gemini\antigravity\scratch\muressons-sim"

FACILITATOR_MANUAL = os.path.join(
    BASE_DIR, "Muressons_Facilitator_Manual_v4_with_Constellation.docx"
)

# ─────────────────────────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────────────────────────
def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)

def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.style.font.size = Pt(11)
    return p

def add_bullet(doc, text):
    doc.add_paragraph(text, style='List Bullet')

def add_table_row(table, *cols):
    row = table.add_row()
    for i, val in enumerate(cols):
        if i < len(row.cells):
            row.cells[i].text = val

def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Cm(1)
    return p


# ─────────────────────────────────────────────────────────────────
#  APPENDIX CONTENT
# ─────────────────────────────────────────────────────────────────
def add_factory_reset_appendix(doc):
    doc.add_page_break()

    # ── Title ──
    add_heading(doc, 'Appendix: Factory Reset & Data Deletion Reference', level=1)

    add_body(doc,
        "This appendix documents exactly what the \"Factory Reset\" command in God Mode "
        "destroys when a Super Admin types DELETE ALL DATA and confirms the nuclear wipe. "
        "Understanding the scope of this action is critical before using it in a live "
        "teaching environment."
    )

    # ── Overview ──
    add_heading(doc, 'Overview', level=2)
    add_body(doc,
        "The Factory Reset is a two-stage safeguarded operation accessible only from "
        "God Mode → Danger Zone → Factory Reset. It requires the Super Admin to:"
    )
    add_bullet(doc, "Type the exact phrase DELETE ALL DATA into a confirmation field.")
    add_bullet(doc, "Click the ☢️ Factory Nuke All Sessions button.")
    add_bullet(doc, "Wait through a mandatory 15-second countdown (which can be cancelled).")
    add_body(doc,
        "Once the countdown reaches zero, the system calls the backend endpoint "
        "DELETE /api/admin/reset-all, which performs a hard delete across all "
        "data stores. This action cannot be undone."
    )

    # ═══════════════════════════════════════════════════════════════
    #  SECTION 1 — WHAT IS DELETED
    # ═══════════════════════════════════════════════════════════════
    add_heading(doc, 'What the Factory Reset Deletes', level=2)

    add_body(doc,
        "The reset operates across two layers: the database persistence layer and the "
        "admin application layer. The following tables itemise every data store that is "
        "wiped."
    )

    # ── Database Layer ──
    add_heading(doc, '1. Database Layer (Persistent Storage)', level=3)
    add_body(doc,
        "These four core data stores are completely cleared. Their contents are also "
        "removed from the on-disk JSON persistence file, meaning the data cannot be "
        "recovered after a server restart."
    )

    tbl_db = doc.add_table(rows=1, cols=3)
    tbl_db.style = 'Table Grid'
    hdr = tbl_db.rows[0].cells
    hdr[0].text = 'Data Store'
    hdr[1].text = 'Contents'
    hdr[2].text = 'Impact of Deletion'

    db_rows = [
        ('Sessions Registry\n(_sessions)',
         'All session metadata: cohort names, facilitator assignments, '
         'player IDs, parent–child relationships, paradigm settings, '
         'currency, difficulty tier, short codes (e.g. SIM-A3K7), '
         'pedagogical overrides, pacing mode, shuffle seeds, start/end dates, '
         'decade forward plans.',
         'Every cohort and every player account vanishes. No session can be '
         'loaded or resumed. All join codes become invalid.'),

        ('Global State History\n(_global_states)',
         'Every round\'s global state for every session: corporate treasury, '
         'group reputation, synergy multiplier, cost of capital, all active '
         'event flags, EBITDA history, tCO₂e emissions, VRIO capabilities, '
         'stakeholder map data, tipping point state, pending capex projects, '
         'inflation index, SDG metrics (political capital, community trust, '
         'emissions intensity), BU substitution state, and autonomous '
         'stakeholder agent state.',
         'All simulation progress is erased. Round-by-round state history '
         'used by the Executive Cockpit, ESG Constellation, Consequence DNA, '
         'and debrief tools is permanently lost.'),

        ('Business Unit States\n(_bu_states)',
         'Per-round snapshots for every business unit: revenue base, opex base, '
         'natural capital debt, social licence score, reputation score, '
         'governance risk score, water dependency, carbon intensity, staff '
         'burnout index, bed capacity utilisation, patient outcomes score, '
         'SDG cluster scores (basic needs, human capital, sustainable growth, '
         'planet, governance, partnerships), population, migration pressure, '
         'and risk factors.',
         'All BU-level performance history is destroyed. The balanced scorecard, '
         'risk radar, and BU comparison tools will show no data.'),

        ('Decision Audit Log\n(_decision_log)',
         'The complete audit trail of every decision made across all sessions: '
         'decision node ID, choice selected (e.g. option_a), capex allocated, '
         'player ID, time-to-decision in seconds, and team consensus method.',
         'All decision history is lost. The Decision Timeline, Debrief Report, '
         'and Consequence DNA Sankey diagram will have no data to render. '
         'Academic assessment of player choices is no longer possible.'),
    ]
    for store, contents, impact in db_rows:
        add_table_row(tbl_db, store, contents, impact)

    doc.add_paragraph()  # spacer

    # ── Admin Application Layer ──
    add_heading(doc, '2. Admin Application Layer (In-Memory State)', level=3)
    add_body(doc,
        "In addition to the persistent database, several in-memory admin data stores "
        "are cleared during the reset:"
    )

    tbl_admin = doc.add_table(rows=1, cols=2)
    tbl_admin.style = 'Table Grid'
    hdr2 = tbl_admin.rows[0].cells
    hdr2[0].text = 'Data Store'
    hdr2[1].text = 'What Is Lost'

    admin_rows = [
        ('God Mode Audit Log',
         'All logged God Mode actions (e.g. crisis triggers, manual overrides, '
         'facilitator registrations). The Activity Log panel will be empty.'),

        ('Crisis Trigger History',
         'The history of all manually triggered crises (e.g. forced strikes, '
         'carbon tax increases). Future crisis scheduling is also cleared.'),

        ('Session Messages',
         'All facilitator broadcast messages sent to specific sessions '
         'via the Universal Broadcast tool.'),

        ('Session Interventions',
         'All active interventions pushed to player sessions '
         '(e.g. live scenario overrides, custom Black Swan events).'),

        ('Practice Mode Flags',
         'Practice mode on/off state for all sessions. Any sessions '
         'that were in practice mode lose that designation.'),

        ('Round Pacing Configuration',
         'Per-session round gating and pacing config. Any cohorts '
         'using instructor-paced or timed-release modes lose those settings.'),

        ('Facilitator Registry',
         'ALL facilitator accounts are deleted and replaced with a single '
         'default Super Admin facilitator (god_mode). Every custom facilitator '
         'created via the Facilitator Registry — including Lead Facilitators '
         'and Base Facilitators — is permanently removed. Their passwords '
         'and scope assignments are lost.'),

        ('Shared Marketplace State',
         'Carbon Credit Pool: all purchases are cleared and the price history '
         'is reset to the baseline [$50,000 per credit]. '
         'Green Talent Pool: all hires are cleared and the cost history '
         'is reset to the baseline [$200,000 per hire].'),

        ('Facilitator Notes',
         'All per-session facilitator notes and comments are erased.'),

        ('Student Bonuses & Peer Evaluations',
         'Any bonus scores awarded to students and peer evaluation data '
         'are deleted.'),

        ('Session Annotations',
         'Per-session annotations used for tagging and categorising '
         'session behaviour are removed.'),
    ]
    for store, lost in admin_rows:
        add_table_row(tbl_admin, store, lost)

    doc.add_paragraph()  # spacer

    # ═══════════════════════════════════════════════════════════════
    #  SECTION 2 — WHAT IS NOT DELETED
    # ═══════════════════════════════════════════════════════════════
    add_heading(doc, 'What the Factory Reset Does NOT Delete', level=2)

    add_body(doc,
        "The following items survive a Factory Reset and remain intact after the wipe:"
    )

    tbl_safe = doc.add_table(rows=1, cols=2)
    tbl_safe.style = 'Table Grid'
    tbl_safe.rows[0].cells[0].text = 'Item'
    tbl_safe.rows[0].cells[1].text = 'Why It Survives'

    safe_rows = [
        ('God Mode Password',
         'The master God Mode password is stored in a separate persistence '
         'file and is not touched by the reset. You will not be locked out.'),

        ('Resource Library',
         'Uploaded teaching resources (PDFs, links, documents) managed via '
         'the Resource Manager are stored independently and are not cleared.'),

        ('Glossary Entries',
         'Custom glossary definitions added via the Glossary Manager survive the reset.'),

        ('Uploaded Media Files',
         'Physical files uploaded for interventions (audio, video, images) remain '
         'on disk in the /public/uploads/ directory. However, the intervention '
         'records that reference them are cleared.'),

        ('Master Overrides & Swipe Files',
         'The God Mode master intervention database (e.g. Global Macro Shift, '
         'Omni-Tech Poach, Force Strike) and all swipe file presets are NOT '
         'cleared. These are defined at the application level and persist.'),

        ('Round Configuration',
         'The 10-round scenario configuration (crisis definitions, options, '
         'flags, impacts) is static application code and is never modified '
         'by the reset.'),

        ('Application Code & Settings',
         'All frontend components, backend logic, and UI configuration are '
         'unaffected. The simulation is fully functional immediately after '
         'a reset — it simply has no sessions to display.'),
    ]
    for item, reason in safe_rows:
        add_table_row(tbl_safe, item, reason)

    doc.add_paragraph()  # spacer

    # ═══════════════════════════════════════════════════════════════
    #  SECTION 3 — TARGETED ALTERNATIVES
    # ═══════════════════════════════════════════════════════════════
    add_heading(doc, 'Targeted Alternatives to Factory Reset', level=2)
    add_body(doc,
        "Before using the Factory Reset, consider whether one of these less destructive "
        "options in the Danger Zone panel would suffice:"
    )

    tbl_alt = doc.add_table(rows=1, cols=3)
    tbl_alt.style = 'Table Grid'
    tbl_alt.rows[0].cells[0].text = 'Action'
    tbl_alt.rows[0].cells[1].text = 'Scope'
    tbl_alt.rows[0].cells[2].text = 'Use When'

    alt_rows = [
        ('Delete Selected Cohorts',
         'Removes only the checked cohorts and their child player sessions.',
         'You want to clean up specific completed or test cohorts '
         'without affecting other active classes.'),

        ('Clear All Orphans Globally',
         'Removes only cohorts that have no facilitator assigned.',
         'You want to clean up sessions created during testing '
         'or by deleted facilitator accounts.'),

        ('Undo Round (Facilitator Dashboard)',
         'Rolls a specific session back by one or more rounds.',
         'A cohort made a decision in error and needs to retry. '
         'No data is permanently deleted — it can be replayed.'),
    ]
    for action, scope, use_when in alt_rows:
        add_table_row(tbl_alt, action, scope, use_when)

    doc.add_paragraph()  # spacer

    # ═══════════════════════════════════════════════════════════════
    #  SECTION 4 — PRE-RESET CHECKLIST
    # ═══════════════════════════════════════════════════════════════
    add_heading(doc, 'Pre-Reset Checklist', level=2)
    add_body(doc,
        "If you have confirmed that a full Factory Reset is necessary (e.g. between "
        "semesters or after a major pilot), complete the following checklist before "
        "proceeding:"
    )

    add_bullet(doc,
        "☐ Export all session data — Use God Mode → System Export to download "
        "JSON snapshots of every active session before they are destroyed."
    )
    add_bullet(doc,
        "☐ Download debrief reports — Generate and save Debrief Reports for any "
        "cohorts that require academic grading or archival."
    )
    add_bullet(doc,
        "☐ Record facilitator accounts — Note down the names and IDs of all "
        "custom facilitator accounts, as they will need to be recreated after the reset."
    )
    add_bullet(doc,
        "☐ Confirm no active players — Verify that no students are currently "
        "mid-session. Active WebSocket connections will be severed without warning."
    )
    add_bullet(doc,
        "☐ Communicate to teaching team — Inform all facilitators that the reset "
        "is happening, as their dashboard will show zero cohorts afterward."
    )

    # ═══════════════════════════════════════════════════════════════
    #  SECTION 5 — POST-RESET STATE
    # ═══════════════════════════════════════════════════════════════
    add_heading(doc, 'Post-Reset State', level=2)
    add_body(doc,
        "After a successful Factory Reset, the system returns to a clean-install state:"
    )
    add_bullet(doc, "Zero cohorts and zero player sessions.")
    add_bullet(doc, "A single default Super Admin facilitator account (god_mode).")
    add_bullet(doc,
        "The shared marketplace is reset: carbon credits at $50,000/credit "
        "and green talent at $200,000/hire."
    )
    add_bullet(doc, "All admin logs, notes, and intervention history are empty.")
    add_bullet(doc,
        "The God Mode dashboard is fully functional and ready to provision "
        "new cohorts immediately."
    )
    add_body(doc,
        "The confirmation message displayed after a successful reset reads: "
        "\"✅ All [N] sessions wiped.\" where [N] is the total number of sessions "
        "(including player sub-sessions) that were destroyed."
    )


# ─────────────────────────────────────────────────────────────────
#  RUNNER
# ─────────────────────────────────────────────────────────────────
def process(filepath, add_fn, label, output_name):
    if not os.path.exists(filepath):
        print(f"[ERROR] File not found: {filepath}")
        return

    doc = docx.Document(filepath)
    add_fn(doc)

    out = os.path.join(BASE_DIR, output_name)
    doc.save(out)
    print(f"[OK] {label} saved -> {output_name}")


if __name__ == '__main__':
    process(
        FACILITATOR_MANUAL,
        add_factory_reset_appendix,
        "Facilitator Manual",
        "Muressons_Facilitator_Manual_v5.docx",
    )
