"""Bulk player provisioning from Excel — single-cohort and master workbook.

Two shapes, one parser core:

  * FACILITATOR shape — a flat sheet of players for ONE cohort the facilitator
    already owns. The cohort is named by the URL, not the sheet, and its
    SIMULATION MODE decides which columns the sheet has: a 4-BU conglomerate
    roster is name/email/programme only, because every player runs all four
    business units; a single-business roster adds industry_vertical + region_id,
    because there each player runs their own company. That is not a cosmetic
    difference — see roster_shape.py, which owns the distinction and explains
    why a conglomerate row carrying a business unit corrupts the cohort
    silently. Both the template and this parser read the SAME RosterShape, so a
    downloaded template can never fail its own upload.

  * PROJECT-ADMIN master workbook — three sheets, `Facilitators`, `Cohorts`
    and `Players`, that provision an entire programme in one upload. Rows are
    linked by a `ref` the author invents (e.g. "f1", "c1"), NOT by an ID the
    server has not assigned yet: the whole point is that none of these entities
    exist before the upload, so the author cannot cite a FAC-003 that does not
    exist. A cohort may alternatively cite an EXISTING facilitator id, which is
    how a second intake gets added to a facilitator provisioned last term.

Contract (owner decision, July 2026): **all-or-nothing**. The workbook is
parsed and validated in full before anything is created; one bad row aborts
the entire import. This differs deliberately from the older facilitator-only
bulk upload, which creates good rows and reports bad ones. That is defensible
for one flat entity; it is not defensible across three linked ones, because a
partial import can leave a cohort pointing at a facilitator that was never
created, and no operator can tell from the resulting state which half applied.

Every error carries sheet + row so the author can fix the file in one pass
rather than re-uploading to discover the next problem.
"""
from __future__ import annotations

import re
from io import BytesIO

from player_capacity import MAX_PLAYERS_CEILING, clamp_max_players
from roster_shape import (
    ASSIGNED_BU, INDUSTRY, KNOWN_COLUMNS, REGION, RosterShape,
    resolve_roster_shape, shape_from_cohort_row, vertical_slot,
)

MAX_UPLOAD_BYTES = 2 * 1024 * 1024

# The data sheet's name. The parser selects it BY NAME rather than using
# wb.active, because Excel persists the selected tab: a template saved while the
# author was reading the 'How to use' sheet would otherwise parse as "header row
# must contain a name column". The facilitator template already carried this
# fix; the player template did not.
PLAYER_SHEET = "Players"

# ── Column vocabulary ───────────────────────────────────────────────────────
# The superset the parser can RECOGNISE. Which of these a given sheet may
# actually carry is the cohort's RosterShape, not this tuple — recognising a
# column is what lets the parser refuse it with an explanation instead of
# silently ignoring it as institutional bookkeeping.
PLAYER_COLUMNS = ("name", "email", "programme", "assigned_bu", "region_id",
                  "industry_vertical")
PLAYER_REQUIRED = ("name",)

# Authors write human headers; we normalise. Unknown columns are ignored rather
# than rejected so an institution can keep its own bookkeeping columns in the
# same sheet.
_ALIASES = {
    "full_name": "name",
    "student_name": "name",
    "player_name": "name",
    "participant": "name",
    "participant_name": "name",
    "e_mail": "email",
    "email_address": "email",
    "email_id": "email",
    "course": "programme",
    "program": "programme",
    "program_name": "programme",
    "programme_name": "programme",
    "batch": "programme",
    # assigned_bu aliases are kept even though the column is never authored in
    # either mode: normalising them is what lets the parser REFUSE the column by
    # name with a remedy, rather than ignoring "Business Unit" as an unknown
    # bookkeeping column and dropping the author's intent on the floor.
    "business_unit": "assigned_bu",
    "bu": "assigned_bu",
    "assigned_business_unit": "assigned_bu",
    "industry": "industry_vertical",
    "vertical": "industry_vertical",
    "sector": "industry_vertical",
    "industry_sector": "industry_vertical",
    "region": "region_id",
    "market": "region_id",
    "cohort": "cohort_ref",
    "cohort_id": "cohort_ref",
    "cohort_key": "cohort_ref",
    "facilitator": "facilitator_ref",
    "facilitator_id": "facilitator_ref",
    "facilitator_key": "facilitator_ref",
    "cohort_title": "cohort_name",
    "session_name": "cohort_name",
    "max_players": "max_players",
    "seats": "max_players",
}

COHORT_COLUMNS = (
    "ref", "cohort_name", "facilitator_ref", "max_players",
    "simulation_mode", "industry_vertical", "region_id", "decision_paradigm",
)
FACILITATOR_COLUMNS = (
    "ref", "name", "email", "contact_number", "programme", "role", "max_cohorts",
)

MASTER_SHEETS = ("Facilitators", "Cohorts", "Players")

# An email is validated only loosely: this is a classroom roster, and rejecting
# an unusual-but-valid institutional address would be worse than accepting a
# typo the facilitator can see in the preview.
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class BulkPlayerError(ValueError):
    """Carries the full error list so the caller can render all of them."""

    def __init__(self, errors: list[dict]):
        self.errors = errors
        n = len(errors)
        head = "; ".join(
            f"{e.get('sheet', 'Sheet')} row {e.get('row', '?')}: {e['error']}"
            for e in errors[:8]
        )
        more = f" (+{n - 8} more)" if n > 8 else ""
        super().__init__(f"{n} error(s) — nothing was created. {head}{more}")


def _norm_header(cell) -> str:
    h = str(cell or "").strip().lower()
    h = re.sub(r"[^a-z0-9]+", "_", h).strip("_")
    return _ALIASES.get(h, h)


def _load(raw: bytes):
    from openpyxl import load_workbook
    if len(raw) > MAX_UPLOAD_BYTES:
        raise BulkPlayerError([{"sheet": "-", "row": 0, "error": "File too large (max 2 MB)"}])
    try:
        return load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise BulkPlayerError([{
            "sheet": "-", "row": 0,
            "error": "Not a readable .xlsx file. Download the template and try again.",
        }])


def _rows_of(ws, columns: tuple[str, ...], sheet_label: str):
    """Yield (rownum, cell_accessor) for each non-blank data row, plus the
    normalised header so the caller can police which columns are PRESENT — a
    check that cannot be made per row, because a column the author added and
    left empty is still a column they believed in.

    Returns ([], [error], header) when the header row lacks a usable column so
    the caller reports a bad sheet rather than silently importing zero rows — a
    silent zero-row import is the failure mode that makes an operator think the
    upload worked.
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], [{"sheet": sheet_label, "row": 1, "error": "Sheet is empty"}], []
    header = [_norm_header(c) for c in rows[0]]
    idx = {name: header.index(name) for name in columns if name in header}
    if "name" not in idx and "ref" not in idx:
        return [], [{
            "sheet": sheet_label, "row": 1,
            "error": f"Header row must contain a 'name' column. Found: {[h for h in header if h]}",
        }], header

    out = []
    for rownum, row in enumerate(rows[1:], start=2):
        if not any(str(c or "").strip() for c in row):
            continue  # fully blank rows are skipped silently

        def cell(key, default="", _row=row):
            i = idx.get(key)
            v = _row[i] if i is not None and i < len(_row) else None
            if v is None:
                return default
            s = str(v).strip()
            return s if s else default

        out.append((rownum, cell))
    return out, [], header


def _sheet_named(wb, name: str):
    """The named sheet, case-insensitively, falling back to wb.active.

    See PLAYER_SHEET: Excel persists which tab was selected at save time, so
    trusting wb.active makes the shipped template fail its own parser whenever
    the author last looked at the guidance tab.
    """
    for title in wb.sheetnames:
        if title.strip().lower() == name.strip().lower():
            return wb[title]
    return wb.active


def _refusal_errors(header: list[str], shape: RosterShape, sheet_label: str) -> list[dict]:
    """Columns this cohort's shape refuses, reported against the HEADER row.

    Refusing at the header rather than per row matters twice over: the author
    gets one error instead of one per player, and a column added but left blank
    is still caught — the author's belief that the column does something is the
    thing being corrected.
    """
    out = []
    for col in dict.fromkeys(header):          # dedupe, preserve author order
        if not col or col not in KNOWN_COLUMNS:
            continue                            # institutional bookkeeping: ignored
        if shape.allows(col):
            continue
        out.append({"sheet": sheet_label, "row": 1, "error": shape.refusal_for(col)})
    return out


def _validate_scope(cell, column, shape, rownum, sheet_label, errors) -> str | None:
    """Resolve one scope cell (industry / region) against the cohort's own
    vocabulary. Blank inherits the cohort's formation value — see roster_shape
    on why blank is inherit rather than error. Returns None on a bad value so
    the caller drops the row; the error is already recorded."""
    raw = cell(column)
    if not raw:
        return shape.default_for(column)
    allowed = shape.choices.get(column, [])
    if raw in allowed:
        return raw
    # Case / spacing forgiveness only. A near-miss is NOT guessed at: silently
    # resolving 'oil and gas' to 'oil_gas' would make the roster disagree with
    # the dropdown that produced it, and the author would never see it.
    lowered = {a.lower(): a for a in allowed}
    hit = lowered.get(raw.strip().lower().replace(" ", "_").replace("-", "_"))
    if hit:
        return hit
    shown = ", ".join(allowed[:8]) + (" …" if len(allowed) > 8 else "")
    errors.append({
        "sheet": sheet_label, "row": rownum,
        "error": (f"'{raw}' is not a valid {column} for this cohort. "
                  f"Pick from the dropdown: {shown}"),
    })
    return None


def _validate_player(cell, rownum, sheet_label, errors, shape: RosterShape) -> dict | None:
    name = cell("name")
    if not name:
        errors.append({"sheet": sheet_label, "row": rownum, "error": "Missing required 'name'"})
        return None
    email = cell("email")
    if email and not _EMAIL_RE.match(email):
        errors.append({"sheet": sheet_label, "row": rownum, "error": f"'{email}' is not a valid email"})
        return None

    rec = {
        "row": rownum,
        "name": name[:100],
        "email": email[:200],
        # Programme is OPTIONAL by design — an open-enrolment cohort has no
        # programme, and forcing a placeholder would pollute the roster export.
        "programme": cell("programme")[:200],
        "cohort_ref": cell("cohort_ref"),
        # Resolved below. A conglomerate player carries EMPTY scope, not an
        # inherited one: an empty assigned_bu is what keeps join_session from
        # filtering that player down to a single business unit.
        "industry_vertical": "",
        "region_id": "",
        "assigned_bu": "",
    }

    if shape.per_player_scope:
        industry = _validate_scope(cell, INDUSTRY, shape, rownum, sheet_label, errors)
        region = _validate_scope(cell, REGION, shape, rownum, sheet_label, errors)
        if industry is None or region is None:
            return None
        rec["industry_vertical"] = industry[:100]
        rec["region_id"] = region[:50]
        # Derived, never authored — the slot that owns this vertical, resolved
        # through the same map join_session uses.
        rec["assigned_bu"] = vertical_slot(industry)[:100]

    return rec


# ── Form 1: one cohort, one flat sheet ──────────────────────────────────────

def parse_player_sheet(raw: bytes, limit: int = MAX_PLAYERS_CEILING,
                       shape: RosterShape | None = None) -> list[dict]:
    """Players for a single cohort. All-or-nothing.

    `limit` is the REMAINING capacity of the target cohort, not its total: a
    cohort of 20 that already holds 15 accepts 5 more. The check lives here so
    the operator is told the file is too big BEFORE any player is created,
    instead of watching the import die partway at row 16.

    `shape` is the target cohort's RosterShape and decides which columns this
    sheet may carry (see roster_shape.py). It defaults to the CONGLOMERATE shape
    — the system default for `simulation_mode`, and the safe default here,
    because it refuses per-player scope columns rather than accepting values it
    has no cohort to validate them against. Endpoints always pass the real
    shape; the default exists for unit tests and for any caller parsing a sheet
    before a cohort is chosen.
    """
    shape = shape or resolve_roster_shape(None)
    wb = _load(raw)
    ws = _sheet_named(wb, PLAYER_SHEET)
    label = ws.title or PLAYER_SHEET
    rows, errors, header = _rows_of(ws, PLAYER_COLUMNS + ("cohort_ref",), label)
    errors.extend(_refusal_errors(header, shape, label))
    players = []
    seen_emails: dict[str, int] = {}
    seen_names: dict[str, int] = {}
    for rownum, cell in rows:
        rec = _validate_player(cell, rownum, label, errors, shape)
        if rec is None:
            continue
        key = rec["name"].lower()
        if key in seen_names:
            errors.append({"sheet": label, "row": rownum,
                           "error": f"Duplicate player name '{rec['name']}' (also row {seen_names[key]})"})
            continue
        seen_names[key] = rownum
        if rec["email"]:
            ekey = rec["email"].lower()
            if ekey in seen_emails:
                errors.append({"sheet": label, "row": rownum,
                               "error": f"Duplicate email '{rec['email']}' (also row {seen_emails[ekey]})"})
                continue
            seen_emails[ekey] = rownum
        players.append(rec)

    if len(players) > limit:
        errors.append({
            "sheet": label, "row": 0,
            "error": (f"Sheet has {len(players)} players but the cohort has room for "
                      f"{limit}. Remove {len(players) - limit} row(s) or raise the "
                      f"cohort's player limit (max {MAX_PLAYERS_CEILING})."),
        })
    if errors:
        raise BulkPlayerError(errors)
    return players


# ── Form 2: master workbook ─────────────────────────────────────────────────

def parse_master_workbook(raw: bytes) -> dict:
    """Facilitators + Cohorts + Players in one file. All-or-nothing.

    Returns {"facilitators": [...], "cohorts": [...], "players": [...]} with
    every cross-reference already resolved, so the caller creates in order and
    never has to re-check linkage.

    A master workbook provisions cohorts of BOTH modes at once, so its single
    Players sheet cannot have one shape. Each player row is validated against
    the shape of the cohort row it cites — the same resolver the live
    single-cohort upload uses — so a row under a conglomerate cohort is refused
    exactly the columns it would be refused after that cohort exists. Without
    this, the master workbook was a second door to the corruption the
    single-cohort path now blocks.
    """
    wb = _load(raw)
    present = {name.strip().lower(): name for name in wb.sheetnames}
    errors: list[dict] = []

    missing = [s for s in MASTER_SHEETS if s.lower() not in present]
    if missing:
        raise BulkPlayerError([{
            "sheet": "-", "row": 0,
            "error": (f"Master workbook must contain sheets named {', '.join(MASTER_SHEETS)}. "
                      f"Missing: {', '.join(missing)}. Found: {', '.join(wb.sheetnames)}"),
        }])

    # ── Facilitators ────────────────────────────────────────────────────────
    fac_rows, errs, _fac_header = _rows_of(wb[present["facilitators"]], FACILITATOR_COLUMNS, "Facilitators")
    errors.extend(errs)
    facilitators, fac_refs = [], {}
    for rownum, cell in fac_rows:
        name = cell("name")
        if not name:
            errors.append({"sheet": "Facilitators", "row": rownum, "error": "Missing required 'name'"})
            continue
        ref = cell("ref") or name
        rkey = ref.strip().lower()
        if rkey in fac_refs:
            errors.append({"sheet": "Facilitators", "row": rownum,
                           "error": f"Duplicate ref '{ref}' (also row {fac_refs[rkey]['row']})"})
            continue
        email = cell("email")
        if email and not _EMAIL_RE.match(email):
            errors.append({"sheet": "Facilitators", "row": rownum, "error": f"'{email}' is not a valid email"})
            continue
        try:
            max_cohorts = max(1, int(float(cell("max_cohorts", "3") or 3)))
        except (TypeError, ValueError):
            max_cohorts = 3
        rec = {
            "row": rownum, "ref": ref, "name": name[:200], "email": email[:200],
            "contact_number": cell("contact_number")[:50],
            "programme": cell("programme")[:200],
            # Role is caller-scoped at the endpoint (assignable_roles_for); a
            # spreadsheet must never be a route to a role above the uploader's.
            "role": (cell("role", "facilitator") or "facilitator").lower(),
            "max_cohorts": max_cohorts,
        }
        fac_refs[rkey] = rec
        facilitators.append(rec)

    # ── Cohorts ─────────────────────────────────────────────────────────────
    coh_rows, errs, _coh_header = _rows_of(wb[present["cohorts"]], COHORT_COLUMNS, "Cohorts")
    errors.extend(errs)
    cohorts, coh_refs = [], {}
    for rownum, cell in coh_rows:
        cname = cell("cohort_name") or cell("name")
        if not cname:
            errors.append({"sheet": "Cohorts", "row": rownum, "error": "Missing required 'cohort_name'"})
            continue
        ref = cell("ref") or cname
        rkey = ref.strip().lower()
        if rkey in coh_refs:
            errors.append({"sheet": "Cohorts", "row": rownum,
                           "error": f"Duplicate ref '{ref}' (also row {coh_refs[rkey]['row']})"})
            continue
        fref = cell("facilitator_ref")
        if not fref:
            errors.append({"sheet": "Cohorts", "row": rownum,
                           "error": "Missing 'facilitator_ref' — name the ref of a row on the "
                                    "Facilitators sheet, or an existing facilitator id"})
            continue
        rec = {
            "row": rownum, "ref": ref, "cohort_name": cname[:200],
            "facilitator_ref": fref,
            # Resolved below: either a new facilitator from this workbook, or an
            # existing id the endpoint must verify.
            "facilitator_is_new": fref.strip().lower() in fac_refs,
            "max_players": clamp_max_players(cell("max_players") or None),
            "simulation_mode": cell("simulation_mode", "conglomerate") or "conglomerate",
            "industry_vertical": cell("industry_vertical"),
            "region_id": cell("region_id"),
            "decision_paradigm": cell("decision_paradigm", "legacy_abc") or "legacy_abc",
        }
        coh_refs[rkey] = rec
        cohorts.append(rec)

    # ── Players ─────────────────────────────────────────────────────────────
    ply_rows, errs, ply_header = _rows_of(wb[present["players"]], PLAYER_COLUMNS + ("cohort_ref",), "Players")
    errors.extend(errs)

    # One shape per cohort ref, resolved from that cohort's own row. Built up
    # front so a player row is validated against the mode of the cohort it cites
    # rather than against a workbook-wide guess.
    shapes = {rkey: shape_from_cohort_row(rec) for rkey, rec in coh_refs.items()}

    # A header-level refusal cannot be issued here the way it is for a
    # single-cohort sheet: one Players sheet serves cohorts of both modes, so
    # `industry_vertical` is legitimate for some rows and refused for others. The
    # column may exist; the refusal is per row, below, and only for rows whose
    # cohort does not allow it. `assigned_bu` is refused for every mode, so it IS
    # a header-level error and is reported once.
    if ASSIGNED_BU in ply_header:
        errors.append({
            "sheet": "Players", "row": 1,
            "error": shape_from_cohort_row(None).refusal_for(ASSIGNED_BU),
        })

    players = []
    per_cohort: dict[str, list[dict]] = {}
    seen: dict[tuple[str, str], int] = {}
    for rownum, cell in ply_rows:
        # cohort_ref is read BEFORE the row is validated, because the ref decides
        # which shape validates it.
        cref = cell("cohort_ref")
        if not cref:
            errors.append({"sheet": "Players", "row": rownum,
                           "error": "Missing 'cohort_ref' — name the ref of a row on the Cohorts sheet"})
            continue
        ckey = cref.strip().lower()
        if ckey not in coh_refs:
            errors.append({"sheet": "Players", "row": rownum,
                           "error": f"cohort_ref '{cref}' does not match any row on the Cohorts sheet"})
            continue
        row_shape = shapes[ckey]
        rec = _validate_player(cell, rownum, "Players", errors, row_shape)
        if rec is None:
            continue
        # A value in a column this row's cohort does not have is an error, not a
        # silent drop: the author believed it would take effect.
        if not row_shape.per_player_scope:
            for col in (INDUSTRY, REGION):
                if cell(col):
                    errors.append({
                        "sheet": "Players", "row": rownum,
                        "error": (f"'{col}' is set to '{cell(col)}', but cohort '{cref}' is a "
                                  f"4-BU conglomerate — every player there runs all four "
                                  f"business units. Clear the cell for this row."),
                    })
            if any(cell(c) for c in (INDUSTRY, REGION)):
                continue
        # Names and emails must be unique WITHIN a cohort, not globally: the
        # same student legitimately appears in two cohorts of a programme.
        nkey = (ckey, rec["name"].lower())
        if nkey in seen:
            errors.append({"sheet": "Players", "row": rownum,
                           "error": f"Duplicate player '{rec['name']}' in cohort '{cref}' (also row {seen[nkey]})"})
            continue
        seen[nkey] = rownum
        if rec["email"]:
            ekey = (ckey, rec["email"].lower())
            if ekey in seen:
                errors.append({"sheet": "Players", "row": rownum,
                               "error": f"Duplicate email '{rec['email']}' in cohort '{cref}' (also row {seen[ekey]})"})
                continue
            seen[ekey] = rownum
        rec["cohort_ref"] = cref
        per_cohort.setdefault(ckey, []).append(rec)
        players.append(rec)

    # Capacity is checked per cohort against that cohort's own configured cap.
    for ckey, group in per_cohort.items():
        cap = coh_refs[ckey]["max_players"]
        if len(group) > cap:
            errors.append({
                "sheet": "Players", "row": group[cap]["row"],
                "error": (f"Cohort '{coh_refs[ckey]['ref']}' has {len(group)} players but its "
                          f"max_players is {cap} (ceiling {MAX_PLAYERS_CEILING})"),
            })

    # An orphan cohort is a warning-shaped bug: the author almost certainly
    # meant to add players and mistyped a ref, which the ref check above would
    # have flagged only on the Players side.
    for rec in cohorts:
        if not per_cohort.get(rec["ref"].strip().lower()):
            rec["player_count"] = 0
        else:
            rec["player_count"] = len(per_cohort[rec["ref"].strip().lower()])

    if errors:
        raise BulkPlayerError(errors)
    return {"facilitators": facilitators, "cohorts": cohorts, "players": players}


# ── Templates ───────────────────────────────────────────────────────────────

def _style_header(ws, headers, width=22):
    from openpyxl.styles import Font, PatternFill
    ws.append(list(headers))
    fill = PatternFill("solid", fgColor="1E293B")
    for i, _h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        ws.column_dimensions[c.column_letter].width = width


def build_player_template(out_path, shape: RosterShape | None = None) -> str:
    """Single-cohort player roster template, built from the cohort's RosterShape.

    The shape decides the columns, the dropdown vocabularies, the example rows
    and the guidance sheet — so a 4-BU conglomerate cohort downloads a template
    with no per-player business unit to fill in, and a single-business cohort
    downloads one whose industry/region dropdowns hold exactly the values its own
    parser accepts. The template and the parser read the same object; neither can
    drift from the other.

    A shape of None means "no cohort chosen" and yields the conglomerate
    template, matching the parser's default.
    """
    from openpyxl import Workbook
    import excel_dropdowns as xd

    shape = shape or resolve_roster_shape(None)
    wb = Workbook()
    ws = wb.active
    ws.title = PLAYER_SHEET
    cols = list(shape.columns)
    _style_header(ws, cols)
    for row in shape.example_rows():
        ws.append(row)
    ws.freeze_panes = "A2"

    # Only the shape's own choice map — never the module-level catalogue. A
    # dropdown offering a value this cohort's parser rejects is the exact drift
    # excel_dropdowns exists to prevent; a dropdown on a column this cohort does
    # not have is the same bug wearing a hat.
    if shape.choices:
        xd.add_dropdowns(wb, ws, cols, shape.choices)

    notes = wb.create_sheet("How to use")
    for line in shape.help_lines():
        notes.append([line])
    notes.column_dimensions["A"].width = 95
    wb.save(str(out_path))
    return str(out_path)


def build_master_template(out_path) -> str:
    """Three-sheet provisioning template for project admins, with dropdowns on
    every constrained column across all three sheets."""
    from openpyxl import Workbook
    import excel_dropdowns as xd
    import roster_shape as rs

    wb = Workbook()

    fac = wb.active
    fac.title = "Facilitators"
    _style_header(fac, list(FACILITATOR_COLUMNS))
    fac.append(["f1", "Dr Anita Rao", "anita@example.edu", "+91-99999-00000", "MBA", "facilitator", 3])
    fac.append(["f2", "Prof Liam Byrne", "liam@example.edu", "", "EMBA", "lead_facilitator", 5])
    fac.freeze_panes = "A2"

    coh = wb.create_sheet("Cohorts")
    _style_header(coh, list(COHORT_COLUMNS))
    coh.append(["c1", "MBA 2026 — Section A", "f1", 20, "conglomerate", "", "", "legacy_abc"])
    coh.append(["c2", "EMBA Pharma Intensive", "f2", 12, "single_bu", "pharma", "south_asia", "multi_toggles"])
    coh.freeze_panes = "A2"

    ply = wb.create_sheet("Players")
    # No assigned_bu column: the business-unit slot is derived from
    # industry_vertical for single-business cohorts and does not exist at all for
    # conglomerate ones. See roster_shape.py.
    #
    # industry_vertical / region_id belong to rows whose cohort is single_bu.
    # Rows under a conglomerate cohort (c1) leave them BLANK, and the example
    # rows below demonstrate exactly that — a template must teach the contract
    # its own parser enforces, and c1 rows carrying a vertical are refused.
    player_cols = ["cohort_ref", "name", "email", "programme",
                   "industry_vertical", "region_id"]
    _style_header(ply, player_cols)
    ply.append(["c1", "Priya Raman", "priya@example.edu", "MBA 2026", "", ""])
    ply.append(["c1", "Sam Okoye", "sam@example.edu", "MBA 2026", "", ""])
    # c2 is single_bu: two players, two different companies in two markets.
    ply.append(["c2", "Maria Silva", "maria@example.edu", "", "pharma", "south_asia"])
    ply.append(["c2", "Tomas Weber", "tomas@example.edu", "", "semiconductor", "europe"])
    ply.freeze_panes = "A2"

    # Dropdowns per sheet. add_dropdowns appends its option lists to the shared
    # hidden _Options sheet, so three calls do not overwrite each other.
    xd.add_dropdowns(wb, fac, list(FACILITATOR_COLUMNS), xd.master_facilitator_column_choices())
    xd.add_numeric_validation(fac, list(FACILITATOR_COLUMNS), "max_cohorts", 1, 100)
    xd.add_dropdowns(wb, coh, list(COHORT_COLUMNS), xd.master_cohort_column_choices())
    xd.add_numeric_validation(coh, list(COHORT_COLUMNS), "max_players", 1, MAX_PLAYERS_CEILING)
    # The Players sheet serves cohorts of BOTH modes, so its dropdowns offer the
    # full formation vocabulary; which rows may use them is decided per row by
    # the cited cohort's mode at parse time.
    xd.add_dropdowns(wb, ply, player_cols, {
        INDUSTRY: rs.formation_verticals(),
        REGION: rs.formation_regions(),
    })

    notes = wb.create_sheet("How to use")
    for line in (
        ["This ONE workbook provisions facilitators, cohorts and players together."],
        [""],
        ["Linking: invent your own 'ref' values (f1, c1, ...) and cite them."],
        ["  Cohorts.facilitator_ref  -> a Facilitators.ref, OR an existing facilitator id (FAC-003)."],
        ["  Players.cohort_ref       -> a Cohorts.ref."],
        ["Refs are yours to choose because none of these records exist yet."],
        [""],
        [f"max_players is per cohort, default {MAX_PLAYERS_CEILING}, ceiling {MAX_PLAYERS_CEILING}."],
        ["'programme' is optional everywhere. 'name' is always required."],
        [""],
        ["ALL-OR-NOTHING: the whole workbook is validated first. One bad row and"],
        ["nothing at all is created, so a failed upload never leaves half a programme."],
        ["Use Preview to see every error at once before committing."],
        [""],
        ["A role you are not permitted to grant is downgraded to 'facilitator'."],
        [""],
        ["PLAYERS AND THEIR COMPANIES — this depends on the cohort's mode."],
        ["  A cohort with simulation_mode = conglomerate runs every one of its"],
        ["  players across all four business units. Leave industry_vertical and"],
        ["  region_id BLANK for those rows; a value there is refused, because"],
        ["  honouring it would quietly scope that player to a single business."],
        ["  A cohort with simulation_mode = single_bu gives each player their own"],
        ["  company: set industry_vertical and region_id per row, and two players"],
        ["  in the same cohort may differ. Blank inherits the cohort's own values."],
        ["There is no assigned_bu column: the business-unit slot is derived from"],
        ["industry_vertical, so the two can never contradict each other."],
        [""],
        ["DROPDOWNS: role, simulation_mode, industry_vertical, decision_paradigm"],
        ["and region_id are pick-lists — free text is rejected."],
        ["  industry_vertical: " + ", ".join(rs.formation_verticals())],
        ["  region_id: " + ", ".join(rs.formation_regions())],
    ):
        notes.append(line)
    notes.column_dimensions["A"].width = 95
    wb.save(str(out_path))
    return str(out_path)
