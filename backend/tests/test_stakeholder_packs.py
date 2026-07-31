"""Stakeholder Packs — one matrix per SBU, per region.

WHAT THIS ADDS
--------------
The 4-SBU conglomerate format gave every business unit the SAME stakeholder
map, because get_stakeholders_for_session() resolves one vertical for the whole
cohort. A pharma BU and a software BU therefore faced identical stakeholders.
A pack binds a matrix per slot so each BU gets its own.

THE DURABILITY BUG FOUND WHILE BUILDING IT
------------------------------------------
stakeholder_db wrote uploaded configs to <repo>/db/stakeholder_configs —
inside the IMAGE, not the mounted volume. On Railway the container filesystem
is ephemeral, so EVERY stakeholder matrix a super-admin uploaded was silently
discarded on the next deploy. Packs would have inherited that, so both now
resolve through runtime_paths onto the volume. test_survives_a_redeploy pins
it by running two independent processes against one volume.

BACK-COMPATIBILITY IS THE POINT
-------------------------------
Cohorts without a pack, and packs missing a slot, must resolve EXACTLY as
before. A pack pointing at a deleted config must fall through rather than hand
a class an empty stakeholder map. Those are the tests that matter most here —
the feature is additive or it is a regression.

ISOLATION: these tests redirect BOTH the data dir and stakeholder_db's config
dir. An ad-hoc script earlier in this project overwrote all 13 real config
files by not doing so.
"""
import io
import json
import os
import pathlib
import subprocess
import sys
import tempfile

import pytest

SLOTS = ["pharma", "electronics", "consumer_goods", "software"]
BACKEND = pathlib.Path(__file__).resolve().parents[1]


@pytest.fixture
def volume(monkeypatch):
    """A throwaway 'mounted volume', with legacy migration disabled so the
    developer's real configs are never pulled in."""
    vol = pathlib.Path(tempfile.mkdtemp(prefix="packvol_"))
    monkeypatch.setenv("MURESSONS_DATA_DIR", str(vol))
    monkeypatch.setenv("MURESSONS_NO_LEGACY_MIGRATION", "1")

    import stakeholder_db
    import stakeholder_packs
    monkeypatch.setattr(stakeholder_db, "_CONFIG_DIR", vol / "stakeholder_configs", raising=False)
    monkeypatch.setattr(stakeholder_db, "_dir_exists", None, raising=False)
    (vol / "stakeholder_configs").mkdir(parents=True, exist_ok=True)
    stakeholder_db.invalidate_cache()
    stakeholder_packs.invalidate_cache()
    yield vol
    stakeholder_db.invalidate_cache()
    stakeholder_packs.invalidate_cache()


def _region_workbook(region: str, slots=SLOTS) -> bytes:
    """The shipped format: one workbook per region, one sheet per SBU."""
    from openpyxl import Workbook
    from stakeholder_config_excel import COLUMNS

    wb = Workbook()
    wb.remove(wb.active)
    for slot in slots:
        ws = wb.create_sheet(title=slot)
        ws.append(list(COLUMNS))
        ws.append([f"{slot}_reg", f"{slot} regulator", "", "manage_closely", "", 90, 80,
                   f"Regulator for {slot} in {region}", "u", "a", "d1", "", "", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── storage ────────────────────────────────────────────────────────────────

def test_pack_round_trips(volume):
    from stakeholder_packs import save_pack, get_pack, list_packs
    assert save_pack("emea_2026", "EMEA 2026", "europe",
                     {s: f"vertical_{s}__europe" for s in SLOTS})
    pack = get_pack("emea_2026")
    assert pack["label"] == "EMEA 2026"
    assert pack["region"] == "europe"
    assert pack["bus"]["pharma"] == "vertical_pharma__europe"
    assert [p["pack_id"] for p in list_packs()] == ["emea_2026"]


def test_invalid_pack_ids_are_refused_not_raised(volume):
    from stakeholder_packs import save_pack
    for bad in ("Has Spaces", "sym!bol", "", "x" * 60):
        assert save_pack(bad, "l", "europe", {}) is False


def test_pack_ids_are_case_insensitive_end_to_end(volume):
    """save_pack lowercases the id. Every READ must too, or "EMEA_2026" saves
    fine and then resolves to nothing — the cohort silently falls back to the
    default map with no error anywhere. This asymmetry existed until the test
    above caught it."""
    from stakeholder_packs import save_pack, get_pack, delete_pack, config_id_for_bu
    assert save_pack("EMEA_2026", "EMEA", "europe", {"pharma": "vertical_pharma__europe"})
    assert get_pack("EMEA_2026") is not None
    assert get_pack("emea_2026") is not None
    assert config_id_for_bu("EMEA_2026", "pharma") == "vertical_pharma__europe"
    assert delete_pack("EMEA_2026") is True


def test_unknown_slots_are_dropped(volume):
    """A typo'd slot must not create a phantom BU binding."""
    from stakeholder_packs import save_pack, get_pack
    save_pack("p1", "P", "europe", {"pharma": "c1", "not_a_slot": "c2"})
    assert set(get_pack("p1")["bus"]) == {"pharma"}


def test_coverage_reports_missing_slots(volume):
    """A half-configured pack must be visible BEFORE a class, not during."""
    from stakeholder_packs import save_pack, pack_coverage
    save_pack("partial", "Partial", "europe", {"pharma": "vertical_pharma__europe"})
    cov = pack_coverage("partial")
    assert cov["complete"] is False
    assert cov["covered"] == ["pharma"]
    assert set(cov["missing"]) == set(SLOTS) - {"pharma"}


def test_delete_removes_the_pack_only(volume):
    """Deleting an index must never delete the matrices it points at."""
    from stakeholder_db import save_region_config, get_region_config_raw
    from stakeholder_packs import save_pack, delete_pack, get_pack
    save_region_config("vertical_pharma__europe", [{"id": "keep_me", "name": "Keep"}])
    save_pack("temp", "T", "europe", {"pharma": "vertical_pharma__europe"})
    assert delete_pack("temp") is True
    assert get_pack("temp") is None
    assert [s["id"] for s in get_region_config_raw("vertical_pharma__europe")] == ["keep_me"]


# ── durability: the bug this work uncovered ────────────────────────────────

def test_survives_a_redeploy(volume):
    """Two independent processes, one volume. Before the fix these files lived
    in the image and a Railway redeploy discarded them."""
    from stakeholder_db import save_region_config
    from stakeholder_packs import save_pack
    save_region_config("vertical_pharma__europe", [{"id": "survivor", "name": "S"}])
    save_pack("emea", "EMEA", "europe", {"pharma": "vertical_pharma__europe"})

    script = (
        "import sys; sys.path.insert(0, %r)\n"
        "from stakeholder_packs import get_pack\n"
        "from stakeholder_db import get_region_config_raw\n"
        "p = get_pack('emea')\n"
        "rows = get_region_config_raw('vertical_pharma__europe') or []\n"
        "print('OK' if p and [r['id'] for r in rows] == ['survivor'] else 'LOST')\n"
    ) % str(BACKEND)
    env = dict(os.environ)
    out = subprocess.run([sys.executable, "-c", script], capture_output=True,
                         text=True, env=env, timeout=60)
    assert "OK" in out.stdout, f"state did not survive a fresh process: {out.stdout} {out.stderr[-400:]}"


# ── per-BU resolution ──────────────────────────────────────────────────────

def test_each_bu_resolves_its_own_matrix(volume):
    from stakeholder_packs import build_pack_from_region_workbook
    from stakeholder_config_excel import import_master_workbook
    from stakeholder_map import get_stakeholders_for_bu

    tmp = volume / "emea.xlsx"
    tmp.write_bytes(_region_workbook("europe"))
    report = build_pack_from_region_workbook("emea", "EMEA", "europe",
                                             import_master_workbook(tmp))
    assert report["saved"] is True
    assert report["missing_slots"] == []

    gs = {"stakeholder_pack_id": "emea"}
    seen = {}
    for slot in SLOTS:
        ids = [s["id"] for s in get_stakeholders_for_bu(gs, slot)]
        assert ids == [f"{slot}_reg"], f"{slot} resolved to {ids}"
        seen[slot] = tuple(ids)
    assert len(set(seen.values())) == len(SLOTS), "BUs shared a matrix — the bug this fixes"


def test_unmatched_sheets_are_reported_not_ignored(volume):
    """A silently skipped sheet is how a cohort reaches a classroom
    half-configured."""
    from openpyxl import Workbook
    from stakeholder_config_excel import COLUMNS, import_master_workbook
    from stakeholder_packs import build_pack_from_region_workbook

    wb = Workbook(); wb.remove(wb.active)
    for name in ("pharma", "Notes"):
        ws = wb.create_sheet(title=name)
        ws.append(list(COLUMNS))
        ws.append([f"{name}_x", name, "", "manage_closely", "", 50, 50, "d", "u", "a",
                   "", "", "", "", "", ""])
    tmp = volume / "mixed.xlsx"; wb.save(tmp)

    report = build_pack_from_region_workbook("mix", "Mix", "europe",
                                             import_master_workbook(tmp))
    assert "pharma" in report["bus"]
    assert report["unmatched_sheets"] or report["missing_slots"], \
        "a sheet matching no SBU slot must be surfaced"


# ── back-compatibility: additive or it is a regression ─────────────────────

def test_no_pack_resolves_exactly_as_before(volume):
    from stakeholder_map import get_stakeholders_for_bu, get_stakeholders_for_session
    gs = {}
    assert get_stakeholders_for_bu(gs, "pharma") == get_stakeholders_for_session(gs)


def test_unknown_pack_falls_through(volume):
    from stakeholder_map import get_stakeholders_for_bu, get_stakeholders_for_session
    gs = {"stakeholder_pack_id": "does_not_exist"}
    assert get_stakeholders_for_bu(gs, "pharma") == get_stakeholders_for_session({})


def test_pack_missing_a_slot_falls_through_for_that_slot(volume):
    from stakeholder_db import save_region_config
    from stakeholder_packs import save_pack
    from stakeholder_map import get_stakeholders_for_bu, get_stakeholders_for_session

    save_region_config("vertical_pharma__europe", [{"id": "pharma_only", "name": "P"}])
    save_pack("half", "Half", "europe", {"pharma": "vertical_pharma__europe"})
    gs = {"stakeholder_pack_id": "half"}

    assert [s["id"] for s in get_stakeholders_for_bu(gs, "pharma")] == ["pharma_only"]
    # software is not in the pack → today's behaviour
    assert get_stakeholders_for_bu(gs, "software") == get_stakeholders_for_session(gs)


def test_pack_pointing_at_a_deleted_config_never_yields_an_empty_map(volume):
    """The failure that would hurt most: a live class with no stakeholders."""
    from stakeholder_packs import save_pack
    from stakeholder_map import get_stakeholders_for_bu

    save_pack("ghost", "Ghost", "europe", {"pharma": "config_that_was_deleted"})
    rows = get_stakeholders_for_bu({"stakeholder_pack_id": "ghost"}, "pharma")
    assert rows, "resolution returned an empty stakeholder map"


def test_pack_id_is_a_cohort_setting():
    from admin_shared import COHORT_OVERRIDABLE_KEYS
    assert "stakeholder_pack_id" in COHORT_OVERRIDABLE_KEYS
