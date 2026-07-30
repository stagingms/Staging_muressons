"""Bulk player provisioning — capacity, single-cohort sheet, master workbook.

Contract under test:
  * A roster's COLUMNS come from the cohort's simulation mode, not from a fixed
    header. A 4-BU conglomerate roster has no per-player business unit, industry
    or region — naming one is refused, because honouring it would scope that
    player to a single BU and quietly stop the cohort being a conglomerate. A
    single-business roster carries industry_vertical + region_id per player,
    because there each player runs their own company. See roster_shape.py.
  * `assigned_bu` is never authored in either mode: it is DERIVED from the
    industry, so a sheet cannot contain a slot that contradicts its vertical.
  * The roster cap is ONE resolver with a hard ceiling; a configured value
    above the ceiling can never be honoured, whatever wrote it.
  * Both upload shapes are ALL-OR-NOTHING: any error and the parser raises
    with every error listed, so nothing is created and the author fixes the
    file in one pass.
  * `programme` is optional; `name` never is.
  * Cross-sheet refs are resolved by the author's own keys, because none of
    the referenced records exist before the upload.
"""
import pathlib
import tempfile

import pytest
from openpyxl import Workbook, load_workbook

from player_capacity import (
    MAX_PLAYERS_CEILING, DEFAULT_MAX_PLAYERS, clamp_max_players,
    resolve_max_players, capacity_error,
)
from player_bulk_excel import (
    BulkPlayerError, parse_player_sheet, parse_master_workbook,
    build_player_template, build_master_template, FACILITATOR_COLUMNS, COHORT_COLUMNS,
)
from roster_shape import resolve_roster_shape

# The conglomerate roster: identity only. This is the DEFAULT shape, so a sheet
# written against this header needs no shape argument.
PLAYER_HEADER = ["name", "email", "programme"]
# The single-business roster: identity plus the player's own company + market.
SB_HEADER = ["name", "email", "programme", "industry_vertical", "region_id"]
# Master workbook Players sheet — serves cohorts of both modes at once.
MASTER_PLAYER_HEADER = ["cohort_ref", "name", "email", "programme",
                        "industry_vertical", "region_id"]

CONGLOMERATE_SHAPE = resolve_roster_shape({"simulation_mode": "conglomerate",
                                           "cohort_name": "Test Cohort"})
SINGLE_BU_SHAPE = resolve_roster_shape({"simulation_mode": "single_bu",
                                        "industry_vertical": "oil_gas",
                                        "region_id": "south_asia",
                                        "cohort_name": "Test Single"})


def _write(wb) -> bytes:
    path = pathlib.Path(tempfile.mkdtemp()) / "wb.xlsx"
    wb.save(str(path))
    return path.read_bytes()


def _player_sheet(rows, header=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Players"
    ws.append(list(header or PLAYER_HEADER))
    for r in rows:
        ws.append(list(r))
    return _write(wb)


def _sb_sheet(rows, header=None):
    return _player_sheet(rows, header=header or SB_HEADER)


def _master(fac_rows, coh_rows, ply_rows):
    wb = Workbook()
    f = wb.active
    f.title = "Facilitators"
    f.append(list(FACILITATOR_COLUMNS))
    for r in fac_rows:
        f.append(list(r))
    c = wb.create_sheet("Cohorts")
    c.append(list(COHORT_COLUMNS))
    for r in coh_rows:
        c.append(list(r))
    p = wb.create_sheet("Players")
    p.append(list(MASTER_PLAYER_HEADER))
    for r in ply_rows:
        p.append(list(r))
    return _write(wb)


# ── Capacity ────────────────────────────────────────────────────────────────

def test_default_is_the_ceiling():
    """"Allow up to 20 players" means an untouched cohort holds 20."""
    assert DEFAULT_MAX_PLAYERS == MAX_PLAYERS_CEILING == 20
    assert resolve_max_players(None) == 20
    assert resolve_max_players({}) == 20


@pytest.mark.parametrize("raw,expected", [
    (12, 12), ("15", 15), (12.0, 12),
    (999, 20),      # above the ceiling is CLAMPED, never honoured
    (0, 1), (-5, 1),  # a zero-seat cohort would be unjoinable while reporting success
    (None, 20), ("", 20), ("abc", 20), (True, 20),  # bool is not a seat count
])
def test_clamp(raw, expected):
    assert clamp_max_players(raw) == expected


def test_a_persisted_over_ceiling_value_is_still_clamped_on_read():
    """Defence in depth: an older build, a hand-edited snapshot or a direct
    settings write must not be able to smuggle a 50-player cohort past the
    enforcement sites."""
    assert resolve_max_players({"max_players": 50}) == MAX_PLAYERS_CEILING


def test_capacity_error_names_the_configured_limit():
    """A facilitator who set 12 must be told 12 — being told 20 reads as a bug."""
    assert "12" in capacity_error(12)
    assert "20" not in capacity_error(12)


def test_settings_normalisation_clamps_on_the_way_in():
    from admin_shared import normalize_advanced_cohort_settings, COHORT_OVERRIDABLE_KEYS
    assert "max_players" in COHORT_OVERRIDABLE_KEYS
    assert normalize_advanced_cohort_settings({"max_players": 99})["max_players"] == 20
    assert normalize_advanced_cohort_settings({"max_players": 8})["max_players"] == 8


# ── Single-cohort sheet ─────────────────────────────────────────────────────

def test_parses_a_conglomerate_roster():
    """The default shape: identity only, and no player is assigned a company."""
    raw = _player_sheet([
        ["Priya Raman", "priya@x.edu", "MBA 2026"],
        ["Sam Okoye", "sam@x.edu", ""],
    ])
    out = parse_player_sheet(raw)
    assert all(p["assigned_bu"] == "" and p["industry_vertical"] == "" for p in out), \
        "a conglomerate player runs all four BUs; any scope here would filter them to one"
    assert [p["name"] for p in out] == ["Priya Raman", "Sam Okoye"]
    assert out[0]["programme"] == "MBA 2026"
    assert out[1]["programme"] == ""     # optional — absent is not an error


def test_only_name_is_required():
    assert len(parse_player_sheet(_player_sheet([["Solo", "", ""]]))) == 1


def test_missing_name_fails_the_whole_import():
    with pytest.raises(BulkPlayerError, match="Missing required 'name'"):
        parse_player_sheet(_player_sheet([["Ok", "a@x.edu", ""], ["", "b@x.edu", ""]]))


def test_blank_rows_are_skipped_silently():
    """Trailing blank rows are an Excel artefact, not an authoring mistake."""
    raw = _player_sheet([["A", "", ""], [None, None, None], ["B", "", ""]])
    assert len(parse_player_sheet(raw)) == 2


def test_malformed_email_is_an_error_not_a_silent_drop():
    with pytest.raises(BulkPlayerError, match="not a valid email"):
        parse_player_sheet(_player_sheet([["A", "not-an-email", ""]]))


@pytest.mark.parametrize("dupe_col,rows", [
    ("name", [["Ana", "a@x.edu", ""], ["ana", "b@x.edu", ""]]),
    ("email", [["Ana", "a@x.edu", ""], ["Bo", "A@X.edu", ""]]),
])
def test_duplicates_within_a_cohort_fail(dupe_col, rows):
    """Case-insensitive: 'Ana' and 'ana' are the same student to a human."""
    with pytest.raises(BulkPlayerError, match="[Dd]uplicate"):
        parse_player_sheet(_player_sheet(rows))


def test_header_aliases_are_tolerated():
    """Registrars export 'Full Name' / 'Email Address' / 'Program'."""
    raw = _player_sheet(
        [["Priya", "p@x.edu", "MBA"]],
        header=["Full Name", "Email Address", "Program"],
    )
    out = parse_player_sheet(raw)
    assert out[0]["name"] == "Priya" and out[0]["programme"] == "MBA"


def test_unknown_columns_are_ignored_not_rejected():
    """An institution's own bookkeeping columns must not break the import."""
    raw = _player_sheet(
        [["Priya", "p@x.edu", "MBA", "ROLL-77"]],
        header=PLAYER_HEADER + ["internal_roll_no"],
    )
    assert parse_player_sheet(raw)[0]["name"] == "Priya"


def test_oversized_sheet_is_rejected_before_anything_is_created():
    """The whole point of the capacity pre-check: the operator is told the file
    is too big instead of watching the import die partway."""
    rows = [[f"P{i}", "", ""] for i in range(25)]
    with pytest.raises(BulkPlayerError) as exc:
        parse_player_sheet(_player_sheet(rows), limit=20)
    assert "25 players" in str(exc.value) and "room for 20" in str(exc.value)


def test_limit_is_remaining_capacity_not_total():
    rows = [[f"P{i}", "", ""] for i in range(6)]
    assert len(parse_player_sheet(_player_sheet(rows), limit=6)) == 6
    with pytest.raises(BulkPlayerError):
        parse_player_sheet(_player_sheet(rows), limit=5)


def test_exactly_twenty_is_allowed():
    rows = [[f"P{i}", f"p{i}@x.edu", ""] for i in range(20)]
    assert len(parse_player_sheet(_player_sheet(rows))) == 20


def test_every_error_is_reported_together():
    """One upload, one fix pass — not one error per re-upload."""
    raw = _player_sheet([
        # Row 2 has data but no name — an authoring mistake, unlike a FULLY
        # blank row, which is an Excel artefact and is skipped silently.
        ["", "nameless@x.edu", ""],
        ["Ana", "bad-email", ""],
        ["Bo", "b@x.edu", ""],
        ["bo", "c@x.edu", ""],
    ])
    with pytest.raises(BulkPlayerError) as exc:
        parse_player_sheet(raw)
    assert len(exc.value.errors) == 3
    assert "3 error(s)" in str(exc.value)


def test_not_an_xlsx_is_a_clean_error():
    with pytest.raises(BulkPlayerError, match="readable .xlsx"):
        parse_player_sheet(b"this is not a spreadsheet")


# ── Master workbook ─────────────────────────────────────────────────────────

FAC_A = ["f1", "Anita", "anita@x.edu", "", "MBA", "facilitator", 3]
COH_A = ["c1", "MBA Section A", "f1", 20, "conglomerate", "", "", "legacy_abc"]


def test_master_links_three_sheets_by_author_chosen_refs():
    raw = _master(
        [FAC_A, ["f2", "Liam", "liam@x.edu", "", "EMBA", "lead_facilitator", 5]],
        [COH_A, ["c2", "EMBA Pharma", "f2", 12, "single_bu", "pharma", "south_asia", "legacy_abc"]],
        [["c1", "Priya", "p@x.edu", "MBA", "", ""],
         ["c2", "Maria", "m@x.edu", "", "pharma", "south_asia"]],
    )
    out = parse_master_workbook(raw)
    assert len(out["facilitators"]) == 2
    assert len(out["cohorts"]) == 2
    assert [p["cohort_ref"] for p in out["players"]] == ["c1", "c2"]
    assert out["cohorts"][1]["max_players"] == 12


def test_missing_sheet_names_the_missing_sheet():
    wb = Workbook()
    wb.active.title = "Facilitators"
    wb.active.append(list(FACILITATOR_COLUMNS))
    with pytest.raises(BulkPlayerError) as exc:
        parse_master_workbook(_write(wb))
    assert "Cohorts" in str(exc.value) and "Players" in str(exc.value)


def test_sheet_names_are_case_insensitive():
    wb = Workbook()
    f = wb.active
    f.title = "FACILITATORS"
    f.append(list(FACILITATOR_COLUMNS))
    f.append(FAC_A)
    c = wb.create_sheet("cohorts")
    c.append(list(COHORT_COLUMNS))
    c.append(COH_A)
    p = wb.create_sheet("players")
    p.append(list(MASTER_PLAYER_HEADER))
    p.append(["c1", "Priya", "p@x.edu", "", "", ""])
    assert len(parse_master_workbook(_write(wb))["players"]) == 1


def test_player_pointing_at_an_unknown_cohort_is_an_error():
    """Never a silent drop: a typo'd ref must surface, not vanish."""
    raw = _master([FAC_A], [COH_A], [["c9", "Priya", "p@x.edu", "", "", ""]])
    with pytest.raises(BulkPlayerError, match="does not match any row"):
        parse_master_workbook(raw)


def test_cohort_may_cite_an_existing_facilitator_id():
    """The parser accepts an off-workbook ref and flags it as not-new; the
    endpoint checks it against the registry. Parsing must not require server
    state, or it could not be unit-tested at all."""
    raw = _master([FAC_A], [["c1", "Second Intake", "FAC-003", 20, "conglomerate", "", "", "legacy_abc"]], [])
    out = parse_master_workbook(raw)
    assert out["cohorts"][0]["facilitator_is_new"] is False


def test_cohort_max_players_is_clamped():
    raw = _master([FAC_A], [["c1", "Huge", "f1", 500, "conglomerate", "", "", "legacy_abc"]], [])
    assert parse_master_workbook(raw)["cohorts"][0]["max_players"] == MAX_PLAYERS_CEILING


def test_per_cohort_capacity_is_checked_against_that_cohorts_own_cap():
    raw = _master(
        [FAC_A],
        [COH_A, ["c2", "Small", "f1", 2, "conglomerate", "", "", "legacy_abc"]],
        [["c2", f"P{i}", f"p{i}@x.edu", "", "", ""] for i in range(3)],
    )
    with pytest.raises(BulkPlayerError) as exc:
        parse_master_workbook(raw)
    assert "max_players is 2" in str(exc.value)


def test_same_student_may_appear_in_two_cohorts():
    """Uniqueness is per cohort, not global — a student legitimately sits in
    two sections of a programme."""
    raw = _master(
        [FAC_A],
        [COH_A, ["c2", "Section B", "f1", 20, "conglomerate", "", "", "legacy_abc"]],
        [["c1", "Priya", "p@x.edu", "", "", ""], ["c2", "Priya", "p@x.edu", "", "", ""]],
    )
    assert len(parse_master_workbook(raw)["players"]) == 2


def test_duplicate_within_one_cohort_still_fails():
    raw = _master([FAC_A], [COH_A],
                  [["c1", "Priya", "p@x.edu", "", "", ""], ["c1", "priya", "q@x.edu", "", "", ""]])
    with pytest.raises(BulkPlayerError, match="[Dd]uplicate"):
        parse_master_workbook(raw)


def test_duplicate_refs_fail():
    raw = _master([FAC_A, ["f1", "Other", "o@x.edu", "", "", "facilitator", 3]], [COH_A], [])
    with pytest.raises(BulkPlayerError, match="Duplicate ref"):
        parse_master_workbook(raw)


def test_cohort_without_a_facilitator_ref_is_an_error():
    raw = _master([FAC_A], [["c1", "Orphan", "", 20, "conglomerate", "", "", "legacy_abc"]], [])
    with pytest.raises(BulkPlayerError, match="facilitator_ref"):
        parse_master_workbook(raw)


def test_cohort_with_no_players_is_allowed_and_counted():
    """Provisioning an empty cohort is legitimate — players may be added later."""
    out = parse_master_workbook(_master([FAC_A], [COH_A], []))
    assert out["cohorts"][0]["player_count"] == 0


# ── Templates round-trip ────────────────────────────────────────────────────

@pytest.mark.parametrize("shape", [CONGLOMERATE_SHAPE, SINGLE_BU_SHAPE],
                         ids=["conglomerate", "single_bu"])
def test_player_template_round_trips_through_its_own_parser(shape):
    """A template the parser rejects is worse than no template — this is the
    exact defect the stakeholder bulk template shipped with in July 2026. Both
    shapes must hold, since both are now downloadable."""
    out = pathlib.Path(tempfile.mkdtemp()) / "tpl.xlsx"
    build_player_template(out, shape)
    parsed = parse_player_sheet(out.read_bytes(), shape=shape)
    assert len(parsed) == 2
    assert parsed[0]["programme"] and parsed[1]["programme"] == ""


def test_master_template_round_trips_through_its_own_parser():
    out = pathlib.Path(tempfile.mkdtemp()) / "master.xlsx"
    build_master_template(out)
    parsed = parse_master_workbook(out.read_bytes())
    assert parsed["totals"] if "totals" in parsed else True
    assert len(parsed["facilitators"]) == 2
    assert len(parsed["cohorts"]) == 2
    assert len(parsed["players"]) == 4


def test_master_template_guidance_sheet_is_ignored_by_the_parser():
    """The 'How to use' tab has no ID column and must not break the import."""
    out = pathlib.Path(tempfile.mkdtemp()) / "master.xlsx"
    build_master_template(out)
    assert "How to use" in load_workbook(out).sheetnames
    parse_master_workbook(out.read_bytes())  # must not raise


# ── The two roster variants ─────────────────────────────────────────────────
# These pin the behaviour that made this feature necessary. The 4-BU case is a
# CORRECTNESS test, not a validation nicety: join_session reads a per-player
# assigned_bu as Priority 1 and create_session then filters that player's
# bu_states down to it, so a conglomerate roster that carried a business unit
# produced N single-BU players with no error anywhere. Nothing in the old suite
# could tell the difference, because the parser was right and the roster it was
# handed was wrong.

@pytest.mark.parametrize("column,value", [
    ("assigned_bu", "pharma"),
    ("industry_vertical", "oil_gas"),
    ("region_id", "europe"),
])
def test_conglomerate_roster_refuses_a_per_player_scope_column(column, value):
    with pytest.raises(BulkPlayerError) as exc:
        parse_player_sheet(
            _player_sheet([["Ana", "a@x.edu", "MBA", value]],
                          header=PLAYER_HEADER + [column]),
            shape=CONGLOMERATE_SHAPE,
        )
    assert column in str(exc.value)
    assert exc.value.errors[0]["row"] == 1, "refusal belongs to the HEADER row, once"


def test_a_refused_column_is_caught_even_when_every_cell_is_blank():
    """The author added the column believing it would do something. Waiting for a
    non-blank cell would let a whole cohort import 'successfully' and then teach
    the author the column works."""
    with pytest.raises(BulkPlayerError, match="assigned_bu"):
        parse_player_sheet(
            _player_sheet([["Ana", "a@x.edu", "MBA", ""]],
                          header=PLAYER_HEADER + ["assigned_bu"]),
            shape=CONGLOMERATE_SHAPE,
        )


def test_a_refused_column_is_named_through_its_alias():
    """A registrar's sheet says 'Business Unit'. Normalising the header before
    the refusal check is what stops it being silently ignored as one of the
    institution's own bookkeeping columns."""
    with pytest.raises(BulkPlayerError, match="assigned_bu"):
        parse_player_sheet(
            _player_sheet([["Ana", "a@x.edu", "MBA", "pharma"]],
                          header=PLAYER_HEADER + ["Business Unit"]),
            shape=CONGLOMERATE_SHAPE,
        )


def test_single_bu_roster_gives_each_player_their_own_company():
    """The point of the mode: two players, two industries, two markets, two
    different seed slots — all from one sheet."""
    out = parse_player_sheet(_sb_sheet([
        ["Ana", "a@x.edu", "", "oil_gas", "south_asia"],
        ["Bo", "b@x.edu", "", "semiconductor", "europe"],
    ]), shape=SINGLE_BU_SHAPE)
    assert [(p["industry_vertical"], p["region_id"]) for p in out] == [
        ("oil_gas", "south_asia"), ("semiconductor", "europe")]
    # assigned_bu is DERIVED, and the two verticals sit in different slots.
    assert [p["assigned_bu"] for p in out] == ["pharma", "electronics"]


def test_single_bu_refuses_an_authored_slot():
    """industry_vertical=oil_gas + assigned_bu=software has no defensible
    reading, so the column does not exist and the slot is derived."""
    with pytest.raises(BulkPlayerError, match="derived"):
        parse_player_sheet(
            _sb_sheet([["Ana", "a@x.edu", "", "oil_gas", "europe", "software"]],
                      header=SB_HEADER + ["assigned_bu"]),
            shape=SINGLE_BU_SHAPE,
        )


def test_single_bu_blank_scope_inherits_the_cohort():
    """A facilitator who wants one shared market leaves the columns alone."""
    out = parse_player_sheet(_sb_sheet([["Ana", "a@x.edu", "", "", ""]]),
                             shape=SINGLE_BU_SHAPE)
    assert out[0]["industry_vertical"] == "oil_gas"   # the cohort's own vertical
    assert out[0]["region_id"] == "south_asia"
    assert out[0]["assigned_bu"] == "pharma"


def test_single_bu_rejects_a_value_outside_the_formation_catalogue():
    with pytest.raises(BulkPlayerError) as exc:
        parse_player_sheet(_sb_sheet([["Ana", "a@x.edu", "", "crypto_mining", "europe"]]),
                           shape=SINGLE_BU_SHAPE)
    assert "crypto_mining" in str(exc.value)
    # The remedy must list real values, or the author has to go hunting.
    assert "oil_gas" in str(exc.value)


def test_single_bu_forgives_case_and_spacing_but_never_guesses():
    """'Oil Gas' is the same id typed by a human. 'petroleum' is a different
    word, and silently resolving it would make the roster disagree with the
    dropdown that produced it."""
    out = parse_player_sheet(_sb_sheet([["Ana", "a@x.edu", "", "Oil Gas", "North America"]]),
                             shape=SINGLE_BU_SHAPE)
    assert out[0]["industry_vertical"] == "oil_gas"
    assert out[0]["region_id"] == "north_america"
    with pytest.raises(BulkPlayerError):
        parse_player_sheet(_sb_sheet([["Ana", "a@x.edu", "", "petroleum", "europe"]]),
                           shape=SINGLE_BU_SHAPE)


def test_the_region_catalogue_covers_every_region_a_cohort_can_be_formed_in():
    """China was offered by the Create-Cohort form but missing from the roster
    vocabulary, so a cohort formed in China could not name its own region on its
    own roster. Same drift class as the un_sdg paradigm bug, pointing the other
    way."""
    import roster_shape as rs
    from regional_reporting import REGION_FRAMEWORKS
    assert set(rs.formation_regions()) == set(REGION_FRAMEWORKS)
    assert "china" in rs.formation_regions()


def test_the_vertical_catalogue_matches_the_map_join_resolves_through():
    """roster_shape derives verticals from bu_profiles (dependency-light, so the
    fail-open path cannot silently collapse to four slots) while join_session
    resolves through router.VERTICAL_SLOT_MAP. The two must agree, or a roster
    could offer an industry that becomes unresolvable at join."""
    import roster_shape as rs
    from router import VERTICAL_SLOT_MAP
    assert set(rs.formation_verticals()) == set(VERTICAL_SLOT_MAP)
    for vertical, slot in VERTICAL_SLOT_MAP.items():
        assert rs.vertical_slot(vertical) == slot


def test_default_shape_is_the_conglomerate():
    """A parser called without a cohort must refuse scope rather than accept
    values it has nothing to validate against."""
    from roster_shape import CONGLOMERATE, resolve_roster_shape
    assert resolve_roster_shape(None).mode == CONGLOMERATE
    assert resolve_roster_shape({}).mode == CONGLOMERATE
    # An overloaded simulation_mode (the Switchboard writes climate values into
    # the same key) must never read as single-business.
    for raw in ("standard", "advanced_climate", "", None, "CONGLOMERATE"):
        assert resolve_roster_shape({"simulation_mode": raw}).mode == CONGLOMERATE
    assert resolve_roster_shape({"simulation_mode": "single_bu"}).mode == "single_bu"


def test_the_template_never_fails_the_hostile_save():
    """Excel persists the selected tab. A template saved while the author was
    reading 'How to use' must still parse — the facilitator template already
    carried this fix; the player template used wb.active and did not."""
    from openpyxl import load_workbook
    out = pathlib.Path(tempfile.mkdtemp()) / "tpl.xlsx"
    build_player_template(out, SINGLE_BU_SHAPE)
    wb = load_workbook(out)
    wb.active = wb.sheetnames.index("How to use")
    wb.save(out)
    assert len(parse_player_sheet(out.read_bytes(), shape=SINGLE_BU_SHAPE)) == 2


# ── Master workbook: one sheet, both modes ──────────────────────────────────

def test_master_holds_each_player_row_to_its_own_cohorts_mode():
    """One Players sheet serves cohorts of both modes, so the refusal is per row:
    a scope value under a conglomerate cohort is wrong even though the identical
    value under the single_bu cohort in the same file is right."""
    raw = _master(
        [FAC_A],
        [COH_A,  # c1 = conglomerate
         ["c2", "Single", "f1", 20, "single_bu", "oil_gas", "south_asia", "legacy_abc"]],
        [["c1", "Ana", "a@x.edu", "", "pharma", "south_asia"],
         ["c2", "Bo", "b@x.edu", "", "semiconductor", "europe"]],
    )
    with pytest.raises(BulkPlayerError) as exc:
        parse_master_workbook(raw)
    msg = str(exc.value)
    assert "Ana" not in msg and "c1" in msg, "the conglomerate row is the one refused"
    assert len(exc.value.errors) == 2, "one error per offending cell (industry + region)"


def test_master_accepts_the_same_file_once_the_conglomerate_rows_are_cleared():
    raw = _master(
        [FAC_A],
        [COH_A,
         ["c2", "Single", "f1", 20, "single_bu", "oil_gas", "south_asia", "legacy_abc"]],
        [["c1", "Ana", "a@x.edu", "", "", ""],
         ["c2", "Bo", "b@x.edu", "", "semiconductor", "europe"],
         ["c2", "Cy", "c@x.edu", "", "", ""]],
    )
    out = parse_master_workbook(raw)
    by_name = {p["name"]: p for p in out["players"]}
    assert by_name["Ana"]["industry_vertical"] == ""          # conglomerate: no company
    assert by_name["Bo"]["industry_vertical"] == "semiconductor"
    assert by_name["Bo"]["assigned_bu"] == "electronics"      # derived
    assert by_name["Cy"]["industry_vertical"] == "oil_gas"    # inherited from c2


def test_master_refuses_an_assigned_bu_column_outright():
    """Refused for every mode, so it is a header error reported once rather than
    once per row."""
    wb = Workbook()
    f = wb.active
    f.title = "Facilitators"
    f.append(list(FACILITATOR_COLUMNS))
    f.append(FAC_A)
    c = wb.create_sheet("Cohorts")
    c.append(list(COHORT_COLUMNS))
    c.append(COH_A)
    p = wb.create_sheet("Players")
    p.append(MASTER_PLAYER_HEADER + ["assigned_bu"])
    p.append(["c1", "Ana", "a@x.edu", "", "", "", "pharma"])
    p.append(["c1", "Bo", "b@x.edu", "", "", "", "software"])
    with pytest.raises(BulkPlayerError) as exc:
        parse_master_workbook(_write(wb))
    assert sum("assigned_bu" in e["error"] for e in exc.value.errors) == 1
