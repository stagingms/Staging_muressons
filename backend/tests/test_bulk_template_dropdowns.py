"""Facilitator bulk-upload template: real Excel dropdowns, and it still parses.

The template used to be a bare header row — every constrained column had to be
retyped from help text, and a typo only surfaced at upload. Now each such
column carries an Excel data-validation dropdown sourced from the
authoritative catalogues.

Two failure modes this pins:

  1. DRIFT — a dropdown offering a value the server rejects. Found for real
     while building this: the help text advertised the `un_sdg` paradigm but
     POST /api/simulations/start rejects it (422), so an import could produce
     a facilitator whose default paradigm could not create a cohort. The
     dropdown is now built from config.VALID_DECISION_PARADIGMS — the same set
     the validator enforces.

  2. SELF-REJECTION — the template failing its own parser. The template gained
     a hidden `_Options` sheet and a visible `Reference` sheet; the parser used
     wb.active, so a file saved while the Reference tab was selected parsed as
     "no name column". The parser now selects the data sheet BY NAME.
"""
import io

import pytest
from openpyxl import load_workbook


@pytest.fixture(scope="module")
def template_bytes():
    """Build via the real endpoint body so the test exercises shipped code."""
    import asyncio
    from admin_router import facilitator_bulk_upload_template

    async def build():
        resp = await facilitator_bulk_upload_template(_guard=None)
        chunks = [c async for c in resp.body_iterator]
        return b"".join(
            c if isinstance(c, bytes) else str(c).encode() for c in chunks
        )

    loop = asyncio.new_event_loop()
    try:
        return loop.run_until_complete(build())
    finally:
        loop.close()


# ── Structure ───────────────────────────────────────────────────────────────

def test_template_has_data_options_and_reference_sheets(template_bytes):
    wb = load_workbook(io.BytesIO(template_bytes))
    assert wb.sheetnames[0] == "Facilitators", "data sheet must be first/active"
    assert "_Options" in wb.sheetnames
    assert "Reference" in wb.sheetnames
    assert wb["_Options"].sheet_state == "hidden", \
        "the dropdown source sheet must not clutter the author's view"


def test_every_constrained_column_has_a_dropdown(template_bytes):
    from admin_router import _BULK_UPLOAD_COLUMNS
    import excel_dropdowns as xd
    from openpyxl.utils import get_column_letter

    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb["Facilitators"]
    covered = set()
    for dv in ws.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            covered.add("".join(ch for ch in rng.split(":")[0] if ch.isalpha()))

    for field in xd.facilitator_column_choices():
        letter = get_column_letter(_BULK_UPLOAD_COLUMNS.index(field) + 1)
        assert letter in covered, f"{field} (column {letter}) has no dropdown"


def test_max_cohorts_is_bounded(template_bytes):
    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb["Facilitators"]
    whole = [dv for dv in ws.data_validations.dataValidation if dv.type == "whole"]
    assert whole, "max_cohorts should reject non-numeric / out-of-range input"


# ── No drift between dropdown and server validation ─────────────────────────

def test_paradigm_dropdown_matches_the_server_validator(template_bytes):
    """THE drift bug. Every offered paradigm must be one create_session accepts."""
    from config import VALID_DECISION_PARADIGMS
    import excel_dropdowns as xd

    offered = set(xd.decision_paradigms())
    assert offered == set(VALID_DECISION_PARADIGMS), (
        f"dropdown offers {sorted(offered)} but the server accepts "
        f"{sorted(VALID_DECISION_PARADIGMS)} — an author would import a file "
        "that passes validation and then cannot create a cohort"
    )
    assert "un_sdg" not in offered, (
        "un_sdg is rejected by POST /api/simulations/start (422); offering it "
        "in the template is the exact drift this test exists to prevent"
    )


def test_catalogue_backed_lists_are_derived_not_retyped():
    """Pathways and verticals must come from their catalogues, so a new one
    appears in the template automatically."""
    import excel_dropdowns as xd
    from ending_pathways import ALL_PATHWAY_IDS
    import materiality_db as m

    assert set(xd.ending_pathways()) == set(ALL_PATHWAY_IDS)
    assert set(xd.industry_verticals()) == set(m.get_bu_ids())


def test_dropdown_options_are_written_to_the_options_sheet(template_bytes):
    import excel_dropdowns as xd
    wb = load_workbook(io.BytesIO(template_bytes))
    opts = wb["_Options"]
    written = {c.value for row in opts.iter_rows() for c in row if c.value}
    for field, values in xd.facilitator_column_choices().items():
        for v in values:
            assert v in written, f"{field} option '{v}' missing from _Options"


# ── The template still parses ───────────────────────────────────────────────

def test_template_round_trips_through_its_own_parser(template_bytes):
    from admin_router import _parse_bulk_upload_sheet
    rows, errors = _parse_bulk_upload_sheet(template_bytes)
    assert errors == [], f"the shipped template fails its own parser: {errors}"
    assert len(rows) == 2, "both example rows should parse"
    assert rows[0]["name"] == "Dr. A. Example"
    assert rows[1]["side_tracks"] == ["brsr_ngrbc", "supply_chain"]


def test_parser_finds_the_data_sheet_even_if_another_tab_was_active(template_bytes):
    """Excel stores the selected tab. An author who saves while reading the
    Reference sheet must not get 'header row must contain a name column'."""
    from admin_router import _parse_bulk_upload_sheet

    wb = load_workbook(io.BytesIO(template_bytes))
    wb.active = wb.sheetnames.index("Reference")   # simulate the hostile save
    buf = io.BytesIO()
    wb.save(buf)

    rows, errors = _parse_bulk_upload_sheet(buf.getvalue())
    assert errors == []
    assert len(rows) == 2, "parser must select the Facilitators sheet by name"
