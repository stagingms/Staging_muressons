"""Multi-scope ("master file") stakeholder upload.

Contract:
  * One workbook may carry many verticals/regions, in either shape — scope
    columns on one sheet, or one sheet per scope.
  * Parsing is TOTAL before any write: a single bad row fails the whole import,
    because a half-applied bulk load leaves scopes in an unknown state and the
    operator cannot tell which are stale.
  * Scope labels are matched leniently (case, punctuation, emoji, aliases such
    as "India" → south_asia) but always resolve to a strict slug.
  * A row whose scope cannot be resolved is an ERROR, never a silent drop.
"""
import pathlib
import tempfile

import pytest
from openpyxl import Workbook

from stakeholder_bulk_excel import (
    COLUMNS, parse_bulk_workbook, resolve_scope, resolve_sheet_scope, build_bulk_template,
)

ROW_TAIL = ["👤", "monitor", "", "low", "low", "desc", "why", "", "", "", "", "", "", ""]


def _sheet_row(sid, name="X"):
    return [sid, name] + ROW_TAIL


def _write(wb) -> str:
    path = pathlib.Path(tempfile.mkdtemp()) / "wb.xlsx"
    wb.save(str(path))
    return str(path)


def _form_a(rows):
    """One sheet, Vertical/Region scope columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["Vertical", "Region"] + COLUMNS)
    for r in rows:
        ws.append(list(r))
    return _write(wb)


# ── Scope resolution ────────────────────────────────────────────────────────

@pytest.mark.parametrize("vertical,region,expected", [
    ("", "Europe", "europe"),
    ("", "🇮🇳 India", "south_asia"),          # emoji + alias
    ("", "south asia", "south_asia"),
    ("", "NORTH AMERICA", "north_america"),
    ("Pharma", "", "vertical_pharma"),
    ("pharma / healthcare", "", "vertical_pharma"),
    ("Oil & Gas", "", "vertical_oil_gas"),
    ("Canonical", "", "canonical"),
    ("", "", "canonical"),                      # both blank = default set
])
def test_resolve_scope_is_lenient_but_exact(vertical, region, expected):
    cid, err = resolve_scope(vertical, region)
    assert err is None
    assert cid == expected


def test_industry_and_region_produce_a_composite_scope():
    """Changed 2026-07-19: an industry's stakeholder map genuinely differs by
    market, so naming both is now a COMPOSITE scope rather than an error."""
    cid, err = resolve_scope("Pharma", "Europe")
    assert err is None
    assert cid == "vertical_pharma__europe"


@pytest.mark.parametrize("vertical,region,expected", [
    ("pharma", "india", "vertical_pharma__south_asia"),
    ("Chemicals", "Europe", "vertical_chemicals__europe"),
    ("electronics", "north_america", "vertical_electronics__north_america"),
    ("Oil & Gas", "🇮🇳 India", "vertical_oil_gas__south_asia"),
])
def test_composite_scope_ids(vertical, region, expected):
    assert resolve_scope(vertical, region) == (expected, None)


@pytest.mark.parametrize("sheet,expected", [
    ("vertical_pharma__south_asia", "vertical_pharma__south_asia"),
    ("Pharma - India", "vertical_pharma__south_asia"),
    ("Chemicals - Europe", "vertical_chemicals__europe"),
])
def test_composite_sheet_names(sheet, expected):
    """_norm collapses '__' to '_', so composites are matched by known
    industry/region pairs rather than by the delimiter."""
    assert resolve_sheet_scope(sheet) == (expected, None)


def test_stakeholder_ids_may_contain_hyphens():
    """A stakeholder id is a dict key, never a filename — only config_id needs
    the strict slug rule. Authors must not be forced to rewrite ids like
    'in-pharma-cdsco'."""
    path = _form_a([["", "Europe", *_sheet_row("in-pharma-cdsco")]])
    assert parse_bulk_workbook(path)["europe"][0]["id"] == "in-pharma-cdsco"


@pytest.mark.parametrize("vertical,region", [("", "Atlantis"), ("Widgets", "")])
def test_unknown_scope_is_an_error(vertical, region):
    cid, err = resolve_scope(vertical, region)
    assert cid is None and err


@pytest.mark.parametrize("sheet,expected", [
    ("Canonical", "canonical"),
    ("Region - Europe", "europe"),
    ("Vertical - Pharma", "vertical_pharma"),
    ("europe", "europe"),
    ("vertical_technology", "vertical_technology"),
])
def test_resolve_sheet_scope(sheet, expected):
    cid, err = resolve_sheet_scope(sheet)
    assert err is None and cid == expected


def test_sheet_names_cannot_use_a_colon():
    """Excel forbids ':' in sheet titles, so the documented separator is a dash.
    Guard against guidance drifting back to a colon."""
    wb = Workbook()
    with pytest.raises(ValueError):
        wb.active.title = "Region: Europe"


# ── Form A: scope columns ───────────────────────────────────────────────────

def test_form_a_routes_rows_to_many_scopes():
    path = _form_a([
        ["", "India", *_sheet_row("reg_in", "CDSCO")],
        ["", "india", *_sheet_row("act_in", "Coalition")],
        ["Pharma", "", *_sheet_row("pv", "PV Officer")],
        ["technology", "", *_sheet_row("ai", "Ethics")],
        ["Canonical", "", *_sheet_row("wire", "Wire")],
    ])
    scopes = parse_bulk_workbook(path)
    assert set(scopes) == {"south_asia", "vertical_pharma", "vertical_technology", "canonical"}
    assert len(scopes["south_asia"]) == 2


# ── Form B: sheet per scope ─────────────────────────────────────────────────

def test_form_b_routes_sheets_to_scopes():
    wb = Workbook()
    first = True
    for sheet, sid in (("Region - Europe", "eu"), ("Vertical - Oil & Gas", "og"),
                       ("Canonical", "board"), ("north_america", "na")):
        ws = wb.active if first else wb.create_sheet()
        ws.title = sheet
        first = False
        ws.append(COLUMNS)
        ws.append(_sheet_row(sid))
    scopes = parse_bulk_workbook(_write(wb))
    assert set(scopes) == {"europe", "vertical_oil_gas", "canonical", "north_america"}


def test_legacy_single_sheet_is_treated_as_canonical():
    wb = Workbook()
    ws = wb.active
    ws.title = "Stakeholder Matrix"
    ws.append(COLUMNS)
    ws.append(_sheet_row("only"))
    assert set(parse_bulk_workbook(_write(wb))) == {"canonical"}


def test_non_stakeholder_sheets_are_ignored():
    """A guidance/notes tab has no ID column and must not break the import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Canonical"
    ws.append(COLUMNS)
    ws.append(_sheet_row("a"))
    notes = wb.create_sheet("How to use")
    notes.append(["Instructions"])
    notes.append(["Fill one row per stakeholder"])
    assert set(parse_bulk_workbook(_write(wb))) == {"canonical"}


# ── Validation is total ─────────────────────────────────────────────────────

def test_duplicate_id_within_a_scope_fails():
    path = _form_a([["", "Europe", *_sheet_row("dup")], ["", "Europe", *_sheet_row("dup")]])
    with pytest.raises(ValueError, match="[Dd]uplicate"):
        parse_bulk_workbook(path)


def test_same_id_in_different_scopes_is_fine():
    path = _form_a([["", "Europe", *_sheet_row("shared")], ["", "Africa", *_sheet_row("shared")]])
    scopes = parse_bulk_workbook(path)
    assert scopes["europe"][0]["id"] == scopes["africa"][0]["id"] == "shared"


def test_invalid_quadrant_fails_the_whole_import():
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["Region"] + COLUMNS)
    ws.append(["Europe", "ok_row", "Fine", "👤", "monitor", "", "low", "low", "d", "w", "", "", "", "", "", "", ""])
    ws.append(["Europe", "bad_row", "Bad", "👤", "teleport", "", "low", "low", "d", "w", "", "", "", "", "", "", ""])
    with pytest.raises(ValueError, match="invalid"):
        parse_bulk_workbook(_write(wb))


def test_all_errors_are_reported_together():
    path = _form_a([
        ["", "Atlantis", *_sheet_row("a")],
        ["Widgets", "", *_sheet_row("b")],
        ["", "Europe", *_sheet_row("BAD ID!")],
    ])
    with pytest.raises(ValueError) as exc:
        parse_bulk_workbook(path)
    msg = str(exc.value)
    assert "3 error(s)" in msg
    assert "Atlantis" in msg and "Widgets" in msg and "BAD ID!" in msg


# ── Template ────────────────────────────────────────────────────────────────

def test_template_round_trips_through_the_parser():
    out = pathlib.Path(tempfile.mkdtemp()) / "tpl.xlsx"
    build_bulk_template(out, {
        "canonical": [{"id": "c1", "name": "C", "correct_quadrant": "monitor"}],
        "europe": [{"id": "e1", "name": "E", "correct_quadrant": "keep_informed"}],
        "vertical_pharma": [{"id": "p1", "name": "P", "correct_quadrant": "manage_closely"}],
    })
    scopes = parse_bulk_workbook(out)
    assert set(scopes) == {"canonical", "europe", "vertical_pharma"}
    assert scopes["europe"][0]["correct_quadrant"] == "keep_informed"
