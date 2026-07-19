"""Live HTTP end-to-end for bulk player provisioning.

The unit tests prove the PARSERS are correct. This proves the ENDPOINTS are
wired: guards, ownership, capacity, persistence, and the credential
projection. Several defects in this repo's history were invisible to unit
tests because the parser was right and the route was not — a 200 that created
nothing, a role that resolved from the registry instead of the token, a field
the frontend read that the payload never carried.
"""
import io
import os
import pathlib
import tempfile

import pytest
from openpyxl import Workbook

os.environ.setdefault("USE_MEMORY_DB", "true")
os.environ.setdefault("MURESSONS_DATA_DIR", tempfile.mkdtemp(prefix="pbu_e2e_"))

from fastapi.testclient import TestClient  # noqa: E402

GOD = {"facilitator_id": "god_mode", "password": "sim2026@iim"}
PLAYER_HEADER = ["name", "email", "programme", "assigned_bu", "region_id"]


@pytest.fixture(scope="module")
def client():
    _seed_data_dir()
    from main import app
    import admin_router as ar
    with TestClient(app) as c:
        ar._commit_timestamps.clear() if hasattr(ar, "_commit_timestamps") else None
        r = c.post("/api/admin/facilitators/login", json=GOD)
        assert r.status_code == 200, r.text
        yield c


def _seed_data_dir():
    d = pathlib.Path(os.environ["MURESSONS_DATA_DIR"])
    d.mkdir(parents=True, exist_ok=True)
    (d / "sessions_snapshot.json").write_text("{}", encoding="utf-8")
    (d / "facilitators.json").write_text("[]", encoding="utf-8")


def _xlsx(rows, header=None, title="Players") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(list(header or PLAYER_HEADER))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, url, data: bytes, name="roster.xlsx"):
    return client.post(url, files={"file": (
        name, io.BytesIO(data),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def _new_cohort(client, name="E2E Cohort") -> str:
    r = client.post("/api/simulations/start",
                    json={"cohort_name": name, "facilitator_id": "god_mode"})
    assert r.status_code in (200, 201), r.text
    return str(r.json()["session_id"])


# ── Templates are downloadable and self-consistent ──────────────────────────

def test_both_templates_download_as_xlsx(client):
    for url in ("/api/admin/players/bulk-template", "/api/admin/provisioning/master-template"):
        r = client.get(url)
        assert r.status_code == 200, f"{url} -> {r.status_code} {r.text[:200]}"
        assert r.content[:2] == b"PK", f"{url} did not return a zip/xlsx"
        assert "attachment" in r.headers.get("content-disposition", "")


# ── Single-cohort upload ────────────────────────────────────────────────────

def test_preview_creates_nothing(client):
    sid = _new_cohort(client, "Preview Cohort")
    data = _xlsx([["Priya", "p@x.edu", "MBA", "", ""], ["Sam", "s@x.edu", "", "", ""]])
    r = _upload(client, f"/api/admin/{sid}/players/bulk-preview", data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["ok"] and body["total"] == 2 and body["remaining"] == 20

    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    assert row["registered_players"] == [], "preview must not create players"


def test_bulk_upload_creates_players_with_demographics(client):
    sid = _new_cohort(client, "Upload Cohort")
    data = _xlsx([
        ["Priya Raman", "priya@x.edu", "MBA 2026", "pharma", "south_asia"],
        ["Sam Okoye", "sam@x.edu", "", "chemicals", "europe"],
    ])
    r = _upload(client, f"/api/admin/{sid}/players/bulk-upload", data)
    assert r.status_code == 200, r.text
    created = r.json()["created"]
    assert len(created) == 2
    assert all(c["temp_password"] for c in created), "every player needs a distributable password"
    assert created[0]["temp_password"] != created[1]["temp_password"], "passwords must not be shared"
    assert created[0]["programme"] == "MBA 2026" and created[1]["programme"] == ""

    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    roster = {p["name"]: p for p in row["registered_players"]}
    assert set(roster) == {"Priya Raman", "Sam Okoye"}
    assert roster["Priya Raman"]["email"] == "priya@x.edu"
    assert roster["Priya Raman"]["programme"] == "MBA 2026"


def test_uploaded_player_can_actually_log_in(client):
    """The end that matters. A created row that cannot authenticate is the
    QA-2026-07-16 #4 defect: credentials that lived only in one worker's
    memory and were rejected everywhere else."""
    sid = _new_cohort(client, "Login Cohort")
    r = _upload(client, f"/api/admin/{sid}/players/bulk-upload",
                _xlsx([["Login Tester", "lt@x.edu", "", "", ""]]))
    assert r.status_code == 200, r.text
    cred = r.json()["created"][0]

    # player-login is the real front door (it resolves the cohort from the id);
    # it is also the only response that carries must_change_password.
    login = client.post("/api/simulations/player-login",
                        json={"player_id": cred["player_id"], "password": cred["temp_password"]})
    assert login.status_code == 200, login.text
    assert login.json().get("must_change_password") is True

    bad = client.post("/api/simulations/player-login",
                      json={"player_id": cred["player_id"], "password": "definitely-wrong"})
    assert bad.status_code != 200, "a wrong password must not authenticate"


# ── Password reveal policy ──────────────────────────────────────────────────

def test_temp_password_is_readable_without_a_reset(client):
    """Request 1: the facilitator must be able to read the password at any
    point, not be forced to invalidate a credential they already handed out."""
    sid = _new_cohort(client, "Reveal Cohort")
    r = _upload(client, f"/api/admin/{sid}/players/bulk-upload",
                _xlsx([["Reveal Me", "rm@x.edu", "", "", ""]]))
    cred = r.json()["created"][0]

    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    entry = row["registered_players"][0]
    assert entry["temp_password"] == cred["temp_password"]
    assert entry["must_change_password"] is True


def test_temp_password_disappears_once_the_player_owns_it(client):
    """...and stops being readable the moment it is no longer a temp password."""
    sid = _new_cohort(client, "Erase Cohort")
    r = _upload(client, f"/api/admin/{sid}/players/bulk-upload",
                _xlsx([["Erase Me", "em@x.edu", "", "", ""]]))
    cred = r.json()["created"][0]

    chg = client.post("/api/simulations/change-password", json={
        "player_id": cred["player_id"],
        "old_password": cred["temp_password"],
        "new_password": "MyOwnPassword123",
    })
    assert chg.status_code == 200, chg.text

    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    entry = next(p for p in row["registered_players"] if p["player_id"] == cred["player_id"])
    assert entry["temp_password"] == "", "a personal password must not be recoverable by staff"
    assert entry["must_change_password"] is False


def test_roster_never_ships_a_password_hash(client):
    """The projection is an allow-list precisely so this can be asserted."""
    sid = _new_cohort(client, "Hash Cohort")
    _upload(client, f"/api/admin/{sid}/players/bulk-upload",
            _xlsx([["Hash Check", "hc@x.edu", "", "", ""]]))
    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    for entry in row["registered_players"]:
        assert "password" not in entry
        assert "plaintext_password" not in entry
        assert not any(str(v).startswith("$2b$") for v in entry.values())


# ── Capacity ────────────────────────────────────────────────────────────────

def test_twenty_players_fit_in_one_cohort(client):
    """Request 3, proven end to end rather than at the parser."""
    sid = _new_cohort(client, "Full Cohort")
    rows = [[f"Player {i:02d}", f"p{i}@x.edu", "MBA", "", ""] for i in range(20)]
    r = _upload(client, f"/api/admin/{sid}/players/bulk-upload", _xlsx(rows))
    assert r.status_code == 200, r.text
    assert len(r.json()["created"]) == 20

    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    assert len(row["registered_players"]) == 20
    assert row["max_players"] == 20

    # The 21st is refused by the single-player path too, not just the bulk one.
    r21 = client.post(f"/api/admin/{sid}/generate-player")
    assert r21.status_code == 400
    assert "20" in r21.json()["detail"]


def test_oversized_upload_creates_nothing(client):
    """All-or-nothing: a rejected file must leave the roster untouched."""
    sid = _new_cohort(client, "Overflow Cohort")
    rows = [[f"Over {i:02d}", f"o{i}@x.edu", "", "", ""] for i in range(25)]
    r = _upload(client, f"/api/admin/{sid}/players/bulk-upload", _xlsx(rows))
    assert r.status_code == 400, r.text

    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    assert row["registered_players"] == []


def test_a_bad_row_creates_nothing(client):
    sid = _new_cohort(client, "Badrow Cohort")
    data = _xlsx([["Good", "g@x.edu", "", "", ""], ["Bad", "not-an-email", "", "", ""]])
    r = _upload(client, f"/api/admin/{sid}/players/bulk-upload", data)
    assert r.status_code == 400

    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    assert row["registered_players"] == [], "the valid row must NOT have been created"


def test_lowering_the_cap_is_honoured(client):
    """Per-cohort, up to a maximum of 20 — a cohort set to 3 holds 3."""
    sid = _new_cohort(client, "Small Cohort")
    r = client.patch(f"/api/admin/sessions/{sid}/cohort-settings", json={"max_players": 3})
    assert r.status_code == 200, r.text

    rows = [[f"Small {i}", f"s{i}@x.edu", "", "", ""] for i in range(4)]
    over = _upload(client, f"/api/admin/{sid}/players/bulk-upload", _xlsx(rows))
    assert over.status_code == 400
    assert "3" in str(over.json()["detail"])

    ok = _upload(client, f"/api/admin/{sid}/players/bulk-upload", _xlsx(rows[:3]))
    assert ok.status_code == 200, ok.text
    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    assert row["max_players"] == 3


def test_raising_the_cap_above_the_ceiling_is_clamped_not_honoured(client):
    sid = _new_cohort(client, "Greedy Cohort")
    client.patch(f"/api/admin/sessions/{sid}/cohort-settings", json={"max_players": 500})
    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    row = next(s for s in lb if s["session_id"] == sid)
    assert row["max_players"] == 20


# ── Master workbook ─────────────────────────────────────────────────────────

def _master_bytes(fac_rows, coh_rows, ply_rows) -> bytes:
    from player_bulk_excel import FACILITATOR_COLUMNS, COHORT_COLUMNS
    wb = Workbook()
    f = wb.active
    f.title = "Facilitators"
    f.append(list(FACILITATOR_COLUMNS))
    for r in fac_rows:
        f.append(list(r))
    c = wb.create_sheet("Cohorts")
    c.append(list(COHORT_COLUMNS))
    for r in coh_rows:
        c.append(list(r))
    p = wb.create_sheet("Players")
    p.append(["cohort_ref"] + PLAYER_HEADER)
    for r in ply_rows:
        p.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_master_upload_provisions_all_three_entities(client):
    data = _master_bytes(
        [["mf1", "Master Anita", "ma@x.edu", "", "MBA", "facilitator", 3]],
        [["mc1", "Master Section A", "mf1", 20, "conglomerate", "", "", "legacy_abc"],
         ["mc2", "Master Section B", "mf1", 5, "conglomerate", "", "", "legacy_abc"]],
        [["mc1", "MP One", "mp1@x.edu", "MBA 2026", "", ""],
         ["mc1", "MP Two", "mp2@x.edu", "MBA 2026", "", ""],
         ["mc2", "MP Three", "mp3@x.edu", "", "", ""]],
    )
    prev = _upload(client, "/api/admin/provisioning/master-preview", data, "master.xlsx")
    assert prev.status_code == 200 and prev.json()["ok"], prev.text
    assert prev.json()["totals"] == {"facilitators": 1, "cohorts": 2, "players": 3}

    r = _upload(client, "/api/admin/provisioning/master-upload", data, "master.xlsx")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["totals"] == {"facilitators": 1, "cohorts": 2, "players": 3}

    new_fac = body["facilitators"][0]["facilitator_id"]
    facs = client.get("/api/admin/facilitators").json()
    ids = {f["facilitator_id"] for f in (facs.get("facilitators") or facs)}
    assert new_fac in ids

    lb = client.get("/api/admin/leaderboard").json()["leaderboard"]
    by_id = {s["session_id"]: s for s in lb}
    small = next(c for c in body["cohorts"] if c["cohort_name"] == "Master Section B")
    assert by_id[small["session_id"]]["max_players"] == 5
    assert len(by_id[small["session_id"]]["registered_players"]) == 1


def test_master_upload_is_all_or_nothing(client):
    """A workbook that fails validation must leave NO facilitator, NO cohort
    and NO player behind — the whole reason this path is not row-by-row."""
    before_facs = len(client.get("/api/admin/facilitators").json().get("facilitators", []))
    before_sess = len(client.get("/api/admin/leaderboard").json()["leaderboard"])

    data = _master_bytes(
        [["bf1", "Should Not Exist", "sne@x.edu", "", "", "facilitator", 3]],
        [["bc1", "Should Not Exist Cohort", "bf1", 20, "conglomerate", "", "", "legacy_abc"]],
        [["bc9", "Orphan Player", "op@x.edu", "", "", ""]],  # ref points nowhere
    )
    r = _upload(client, "/api/admin/provisioning/master-upload", data, "bad.xlsx")
    assert r.status_code == 400, r.text

    after_facs = len(client.get("/api/admin/facilitators").json().get("facilitators", []))
    after_sess = len(client.get("/api/admin/leaderboard").json()["leaderboard"])
    assert after_facs == before_facs
    assert after_sess == before_sess


def test_master_upload_rejects_an_unknown_facilitator_ref(client):
    data = _master_bytes(
        [],
        [["xc1", "Dangling", "FAC-DOES-NOT-EXIST", 20, "conglomerate", "", "", "legacy_abc"]],
        [],
    )
    r = _upload(client, "/api/admin/provisioning/master-upload", data, "dangle.xlsx")
    assert r.status_code == 400
    assert "facilitator_ref" in str(r.json()["detail"])
