import docx
import os
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def append_to_briefings():
    # Base path
    base_dir = r"c:\Users\Home\.gemini\antigravity\scratch\muressons-sim"
    source_briefing = os.path.join(base_dir, "Muressons_Simulation_Briefings ver 6.docx")
    source_modules = os.path.join(base_dir, "Muressons_Modules_Full_Stack_Guide.docx")
    
    # Try to open latest briefing
    if not os.path.exists(source_briefing):
        print(f"File not found: {source_briefing}. Checking for other versions...")
        # Fallback to the base one if v6 is missing somehow
        source_briefing = os.path.join(base_dir, "Muressons_Simulation_Briefings.docx")
        
    doc = docx.Document(source_briefing)
    
    # --- 1. Append Section 13.5 ---
    doc.add_page_break()
    sec_title = doc.add_heading('Section 13.5: Regulatory Sandbox & Polycentric Governance Outcomes', level=1)
    
    doc.add_heading('Pedagogical Rationale', level=2)
    doc.add_paragraph(
        "The Regulatory Sandbox (SE-7) represents the pinnacle of complex systems thinking within the Muressons simulation. "
        "Available exclusively in the Expert Tier, this module shifts the locus of control from internal corporate strategy "
        "to macroeconomic policy design. Facilitators and advanced cohorts can design, test, and observe the systemic shocks "
        "of custom ESG regulations on the simulated market."
    )
    
    doc.add_heading('Theoretical Base', level=2)
    p1 = doc.add_paragraph(style='List Bullet')
    p1.add_run("Pigou (1920): ").bold = True
    p1.add_run("Internalization of negative externalities via Pigouvian taxation (e.g., carbon pricing mechanisms).")
    
    p2 = doc.add_paragraph(style='List Bullet')
    p2.add_run("Coase (1960): ").bold = True
    p2.add_run("Property rights and transaction costs, forcing players to negotiate rather than relying strictly on top-down tax mandates.")
    
    p3 = doc.add_paragraph(style='List Bullet')
    p3.add_run("Ostrom (2009): ").bold = True
    p3.add_run("Polycentric governance models, highlighting how localized, multi-layered regulations often outperform monolithic global edicts in managing common-pool resources.")
    
    doc.add_heading('Mechanics & Sample Calculations', level=2)
    doc.add_paragraph(
        "When an exogenous shock is introduced (e.g., a $85/ton carbon tax), the simulation intercepts the main game loop, applying a multiplier penalty against the absolute carbon footprint of every Business Unit. "
        "This creates cascading financial stress:"
    )
    math_p = doc.add_paragraph()
    math_p.add_run("Pigouvian Penalty = BU_Carbon_Emissions * Custom_Tax_Rate_Per_Ton\n").font.name = 'Courier New'
    math_p.add_run("Example: 50,000 tons * $85/ton = $4.25M deduction from Corporate Treasury.").font.name = 'Courier New'
    
    doc.add_paragraph(
        "This forces players to radically alter their leverage points, often pushing them to restructure their business models entirely to avoid bankruptcy in subsequent rounds."
    )
    
    # --- 2. Append the "Document Above" (12 Modules) ---
    doc.add_page_break()
    doc.add_heading('Appendix: New Modules Integration & Full-Stack Pedagogical Framework', level=1)
    
    # Read the content from the previously generated document
    if os.path.exists(source_modules):
        mod_doc = docx.Document(source_modules)
        for para in mod_doc.paragraphs:
            # Recreate the paragraph based on its style to maintain headings
            new_p = doc.add_paragraph(para.text, style=para.style.name)
            # Retain bolding
            for run in para.runs:
                # We already copied text, so adding runs would duplicate it if we aren't careful.
                # Actually, doc.add_paragraph(para.text) copies text as a block. 
                # A safer way to copy runs is:
                pass
        
        # Proper deep copy of paragraphs to maintain formatting
        # Remove the naively added paragraphs from the loop above and re-do:
        
        # Clear the ones we just naively added
        for p in list(doc.paragraphs)[-len(mod_doc.paragraphs):]:
            p._element.getparent().remove(p._element)
            
        for para in mod_doc.paragraphs:
            new_p = doc.add_paragraph(style=para.style.name)
            new_p.alignment = para.alignment
            for run in para.runs:
                new_run = new_p.add_run(run.text)
                new_run.bold = run.bold
                new_run.italic = run.italic
                new_run.font.name = run.font.name
                new_run.font.size = run.font.size
    else:
        doc.add_paragraph("[Error: Could not locate the Muressons_Modules_Full_Stack_Guide.docx to append]")

    output_path = os.path.join(base_dir, "Muressons_Simulation_Briefings ver 7.docx")
    doc.save(output_path)
    print(f"Successfully saved to {output_path}")

if __name__ == '__main__':
    append_to_briefings()
