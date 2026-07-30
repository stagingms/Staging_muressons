"""
Muressons Simulation — Stakeholder Matrix Excel Converter
Bidirectional converter between stakeholder config dicts ↔ stakeholder_matrix.xlsx

Usage (standalone):
    python stakeholder_config_excel.py export [output_path]   # dict → Excel
    python stakeholder_config_excel.py import [input_path]    # Excel → dict

Dev/admin dependency only — openpyxl is lazily imported inside functions so it
is NOT required at simulation runtime.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

# ── Constants ────────────────────────────────────────────────
SHEET_NAME = "Stakeholder Matrix"

COLUMNS = [
    "ID",
    "Name",
    "Icon",
    "Correct Quadrant",
    "Alternate Quadrant",
    "Urgency",
    "Legitimacy",
    "Description",
    "Urgency Rationale",
    "Alternate Rationale",
    "Intel Dossier 1",
    "Intel Dossier 2",
    "Intel Dossier 3",
    "Engagement Tactic 1",
    "Engagement Tactic 2",
    "Engagement Tactic 3",
]

VALID_QUADRANTS = {"manage_closely", "keep_informed", "keep_satisfied", "monitor"}
VALID_URGENCY_LEGITIMACY = {"high", "medium", "low"}

_COL_WIDTHS = [
    20,   # ID
    32,   # Name
    6,    # Icon
    18,   # Correct Quadrant
    18,   # Alternate Quadrant
    10,   # Urgency
    10,   # Legitimacy
    60,   # Description
    50,   # Urgency Rationale
    50,   # Alternate Rationale
    60,   # Intel Dossier 1
    60,   # Intel Dossier 2
    60,   # Intel Dossier 3
    60,   # Engagement Tactic 1
    60,   # Engagement Tactic 2
    60,   # Engagement Tactic 3
]


# ═══════════════════════════════════════════════════════════════
#  EXPORT: list[dict] → Excel
# ═══════════════════════════════════════════════════════════════

def _encode_tactic(tactic: dict) -> str:
    """Encode a single engagement tactic dict to pipe-delimited string.

    Format: ``Label|true/false|Rationale``
    """
    label = tactic.get("label", "")
    correct = "true" if tactic.get("correct") else "false"
    rationale = tactic.get("rationale", "")
    return f"{label}|{correct}|{rationale}"


def export_stakeholders_to_excel(
    stakeholders: list[dict],
    output_path: str | Path,
) -> None:
    """Write a list of stakeholder dicts to a richly formatted Excel file.

    openpyxl is imported lazily so this module can be imported at runtime
    without triggering an ImportError when openpyxl is absent.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    # ── Styles ────────────────────────────────────────────
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_font = Font(name="Calibri", size=10)
    id_font = Font(name="Consolas", size=10)
    thin_border = Border(bottom=Side(style="thin", color="D5D8DC"))
    row_fill_even = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    row_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # ── Header row ────────────────────────────────────────
    for col_idx, (header, width) in enumerate(zip(COLUMNS, _COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(stakeholders) + 1}"

    # ── Data rows ─────────────────────────────────────────
    for row_idx, sh in enumerate(stakeholders, 2):
        intel = sh.get("intel_dossier", [])
        tactics = sh.get("engagement_tactics", [])

        values = [
            sh.get("id", ""),
            sh.get("name", ""),
            sh.get("icon", ""),
            sh.get("correct_quadrant", ""),
            sh.get("alternate_quadrant", ""),
            sh.get("urgency", ""),
            sh.get("legitimacy", ""),
            sh.get("description", ""),
            sh.get("urgency_rationale", ""),
            sh.get("alternate_rationale", ""),
            intel[0] if len(intel) > 0 else "",
            intel[1] if len(intel) > 1 else "",
            intel[2] if len(intel) > 2 else "",
            _encode_tactic(tactics[0]) if len(tactics) > 0 else "",
            _encode_tactic(tactics[1]) if len(tactics) > 1 else "",
            _encode_tactic(tactics[2]) if len(tactics) > 2 else "",
        ]

        fill = row_fill_even if (row_idx % 2 == 0) else row_fill_odd

        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = id_font if col_idx == 1 else data_font
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=(col_idx >= 8))
            cell.border = thin_border

    # ── Save ──────────────────────────────────────────────
    ws.sheet_properties.tabColor = "1B4F72"
    wb.save(str(output_path))


# ═══════════════════════════════════════════════════════════════
#  IMPORT: Excel → list[dict]
# ═══════════════════════════════════════════════════════════════

def _parse_tactic(raw: str) -> dict | None:
    """Parse a pipe-delimited engagement tactic string back to a dict.

    Expected format: ``Label|true/false|Rationale``
    Returns None if the string is empty or unparseable.
    """
    if not raw or not isinstance(raw, str):
        return None
    raw = raw.strip()
    if not raw:
        return None
    parts = raw.split("|", 2)
    if len(parts) < 3:
        return None
    label = parts[0].strip()
    correct = parts[1].strip().lower() in ("true", "1", "yes")
    rationale = parts[2].strip()
    if not label:
        return None
    # Generate a stable tactic id from label
    tactic_id = label.lower().replace(" ", "_")[:40]
    return {
        "id": tactic_id,
        "label": label,
        "correct": correct,
        "rationale": rationale,
    }


def import_stakeholders_from_excel(xlsx_path: str | Path) -> list[dict]:
    """Read a Stakeholder Matrix Excel file and return a validated list of stakeholder dicts.

    openpyxl is imported lazily.

    Raises ``ValueError`` on validation failures with a descriptive message.
    """
    from openpyxl import load_workbook

    wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)

    # Try to find the correct sheet
    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        ws = wb.active

    stakeholders: list[dict] = []
    seen_ids: set[str] = set()
    errors: list[str] = []

    # Read header row to build column map
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        raise ValueError("Excel file has no header row.")

    # Build col_name → index mapping (case-insensitive, strip whitespace)
    col_map: dict[str, int] = {}
    for idx, val in enumerate(header_row):
        if val is not None:
            col_map[str(val).strip().lower()] = idx

    def _get(row_values: tuple, col_name: str) -> str:
        """Safely retrieve a cell value by column header name."""
        idx = col_map.get(col_name.lower())
        if idx is None or idx >= len(row_values):
            return ""
        val = row_values[idx]
        return str(val).strip() if val is not None else ""

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row is None or all(cell is None for cell in row):
            continue

        sid = _get(row, "ID")
        if not sid:
            continue  # skip empty rows

        # ── Validate ID uniqueness ────────────────────────
        if sid in seen_ids:
            errors.append(f"Row {row_num}: Duplicate stakeholder ID '{sid}'.")
        seen_ids.add(sid)

        # ── Parse fields ──────────────────────────────────
        name = _get(row, "Name")
        icon = _get(row, "Icon") or "👤"
        correct_quadrant = _get(row, "Correct Quadrant").lower().replace(" ", "_")
        alternate_quadrant = _get(row, "Alternate Quadrant").lower().replace(" ", "_")
        urgency = _get(row, "Urgency").lower()
        legitimacy = _get(row, "Legitimacy").lower()
        description = _get(row, "Description")
        urgency_rationale = _get(row, "Urgency Rationale")
        alternate_rationale = _get(row, "Alternate Rationale")

        # ── Validate quadrants ────────────────────────────
        if correct_quadrant not in VALID_QUADRANTS:
            errors.append(
                f"Row {row_num} ('{sid}'): Invalid correct_quadrant '{correct_quadrant}'. "
                f"Must be one of: {', '.join(sorted(VALID_QUADRANTS))}"
            )
        if alternate_quadrant and alternate_quadrant not in VALID_QUADRANTS:
            errors.append(
                f"Row {row_num} ('{sid}'): Invalid alternate_quadrant '{alternate_quadrant}'. "
                f"Must be one of: {', '.join(sorted(VALID_QUADRANTS))}"
            )

        # ── Validate urgency / legitimacy ─────────────────
        if urgency and urgency not in VALID_URGENCY_LEGITIMACY:
            errors.append(
                f"Row {row_num} ('{sid}'): Invalid urgency '{urgency}'. "
                f"Must be one of: {', '.join(sorted(VALID_URGENCY_LEGITIMACY))}"
            )
        if legitimacy and legitimacy not in VALID_URGENCY_LEGITIMACY:
            errors.append(
                f"Row {row_num} ('{sid}'): Invalid legitimacy '{legitimacy}'. "
                f"Must be one of: {', '.join(sorted(VALID_URGENCY_LEGITIMACY))}"
            )

        # ── Intel dossier ─────────────────────────────────
        intel_dossier: list[str] = []
        for n in range(1, 4):
            val = _get(row, f"Intel Dossier {n}")
            if val:
                intel_dossier.append(val)

        # ── Engagement tactics ────────────────────────────
        engagement_tactics: list[dict] = []
        for n in range(1, 4):
            val = _get(row, f"Engagement Tactic {n}")
            parsed = _parse_tactic(val)
            if parsed:
                engagement_tactics.append(parsed)

        # ── Build stakeholder dict ────────────────────────
        sh: dict[str, Any] = {
            "id": sid,
            "name": name,
            "icon": icon,
            "correct_quadrant": correct_quadrant,
            "urgency": urgency or "low",
            "legitimacy": legitimacy or "low",
            "description": description,
            "urgency_rationale": urgency_rationale,
            "intel_dossier": intel_dossier,
        }
        if alternate_quadrant:
            sh["alternate_quadrant"] = alternate_quadrant
        if alternate_rationale:
            sh["alternate_rationale"] = alternate_rationale
        if engagement_tactics:
            sh["engagement_tactics"] = engagement_tactics

        stakeholders.append(sh)

    wb.close()

    # ── Final validation ──────────────────────────────────
    if not stakeholders:
        raise ValueError("No stakeholders found in the Excel file.")

    if errors:
        raise ValueError(
            f"Validation failed with {len(errors)} error(s):\n"
            + "\n".join(f"  • {e}" for e in errors)
        )

    return stakeholders


# ═══════════════════════════════════════════════════════════════
#  CLI (standalone usage)
# ═══════════════════════════════════════════════════════════════


# ═══════════════════════════════════════════════════════════════════════════
#  MASTER WORKBOOK — every scope in ONE file
#
#  BUG-2026-07-30: the Stakeholder Configurator's "Master file (all verticals &
#  regions)" tab called three endpoints that were never implemented —
#  /bulk-template, /bulk-preview, /bulk-upload. Downloading showed
#  "Template download failed" (a 404 surfacing as a generic message). The whole
#  master-file mode was frontend-only; single-scope worked because
#  /download/{config_id} does exist.
#
#  Two layouts are accepted on the way back IN, because the UI promises both:
#    * one sheet per scope, sheet name == config_id  (what we WRITE)
#    * a single sheet carrying "Vertical" and/or "Region" columns
#  Excel caps sheet names at 31 chars, and config ids like
#  vertical_banking_financial_services__north_america exceed that, so the
#  written workbook also carries a Scopes index mapping sheet -> real id. The
#  reader prefers that index and only falls back to the sheet title.
# ═══════════════════════════════════════════════════════════════════════════

# Column header -> stakeholder dict key. Mirrors the row order written by
# export_stakeholders_to_excel, so the single-sheet and master exports can
# never disagree about what a column means.
_COLUMN_TO_KEY = {
    "ID": "id",
    "Name": "name",
    "Icon": "icon",
    "Correct Quadrant": "correct_quadrant",
    "Alternate Quadrant": "alternate_quadrant",
    "Urgency": "urgency",
    "Legitimacy": "legitimacy",
    "Description": "description",
    "Urgency Rationale": "urgency_rationale",
    "Alternate Rationale": "alternate_rationale",
    "Intel Dossier 1": "intel_dossier",
    "Intel Dossier 2": "intel_dossier",
    "Intel Dossier 3": "intel_dossier",
    "Engagement Tactic 1": "engagement_tactics",
    "Engagement Tactic 2": "engagement_tactics",
    "Engagement Tactic 3": "engagement_tactics",
}

INDEX_SHEET = "Scopes"
_SHEET_LIMIT = 31


def _safe_sheet_name(config_id: str, taken: set[str]) -> str:
    """Excel-legal, unique, and stable-ish for a human reading tabs."""
    bad = set('[]:*?/\\')
    cleaned = "".join(("_" if ch in bad else ch) for ch in config_id)
    name = cleaned[:_SHEET_LIMIT]
    if name not in taken:
        taken.add(name)
        return name
    for i in range(2, 100):
        suffix = f"~{i}"
        cand = name[: _SHEET_LIMIT - len(suffix)] + suffix
        if cand not in taken:
            taken.add(cand)
            return cand
    raise ValueError(f"could not make a unique sheet name for {config_id!r}")


def export_master_workbook(scopes: dict, output_path) -> dict:
    """Write {config_id: [stakeholder dicts]} to one workbook.

    Returns {config_id: sheet_name} so callers can report what was written.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    index_ws = wb.active
    index_ws.title = INDEX_SHEET

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")

    index_ws.append(["Config ID", "Sheet", "Stakeholders"])
    for c in index_ws[1]:
        c.font = header_font
        c.fill = header_fill
    index_ws.column_dimensions["A"].width = 46
    index_ws.column_dimensions["B"].width = 34
    index_ws.column_dimensions["C"].width = 14

    taken: set[str] = {INDEX_SHEET}
    written: dict = {}

    for config_id in sorted(scopes):
        rows = scopes[config_id] or []
        sheet_name = _safe_sheet_name(str(config_id), taken)
        ws = wb.create_sheet(title=sheet_name)
        ws.append(list(COLUMNS))
        for c in ws[1]:
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for sh in rows:
            ws.append([_cell_for(sh, col) for col in COLUMNS])
        for i, col in enumerate(COLUMNS, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = 40 if "Descr" in col or "Rationale" in col or "Dossier" in col or "Tactic" in col else 20
        ws.freeze_panes = "A2"
        index_ws.append([str(config_id), sheet_name, len(rows)])
        written[str(config_id)] = sheet_name

    from openpyxl.styles import Font as _F
    index_ws["A1"].font = header_font
    wb.save(str(output_path))
    return written


def _cell_for(sh: dict, column: str):
    """Reuse the single-sheet encoding so both exports stay identical."""
    key = _COLUMN_TO_KEY.get(column)
    if key is None:
        return ""
    if column.startswith("Engagement Tactic"):
        idx = int(column[-1]) - 1
        tactics = sh.get("engagement_tactics") or []
        return _encode_tactic(tactics[idx]) if idx < len(tactics) else ""
    if column.startswith("Intel Dossier"):
        idx = int(column[-1]) - 1
        dossier = sh.get("intel_dossier") or []
        return dossier[idx] if idx < len(dossier) else ""
    val = sh.get(key, "")
    return "" if val is None else val


def import_master_workbook(xlsx_path) -> dict:
    """Read a master workbook and return {config_id: [stakeholder dicts]}.

    Accepts either layout described in this section's header. Raises ValueError
    with a descriptive message when neither can be found, so the operator is
    told what shape was expected rather than getting an empty result.
    """
    from openpyxl import load_workbook

    wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)

    # Sheet -> config_id, from the index we write.
    sheet_to_id: dict = {}
    if INDEX_SHEET in wb.sheetnames:
        ws = wb[INDEX_SHEET]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        for row in rows:
            if not row:
                continue
            cid = (row[0] or "").strip() if isinstance(row[0], str) else row[0]
            sheet = (row[1] or "").strip() if len(row) > 1 and isinstance(row[1], str) else None
            if cid and sheet:
                sheet_to_id[sheet] = str(cid)

    out: dict = {}
    for sheet_name in wb.sheetnames:
        if sheet_name == INDEX_SHEET:
            continue
        ws = wb[sheet_name]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header:
            continue
        cols = {str(h).strip(): i for i, h in enumerate(header) if h}
        if "ID" not in cols:
            continue

        has_scope_cols = ("Vertical" in cols) or ("Region" in cols)
        grouped: dict = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            if has_scope_cols:
                vert = row[cols["Vertical"]] if "Vertical" in cols and cols["Vertical"] < len(row) else None
                reg = row[cols["Region"]] if "Region" in cols and cols["Region"] < len(row) else None
                cid = _compose_config_id(vert, reg)
            else:
                cid = sheet_to_id.get(sheet_name, sheet_name)
            grouped.setdefault(cid, []).append(row)

        for cid, raw_rows in grouped.items():
            parsed = _rows_to_stakeholders(cols, raw_rows, cid)
            if parsed:
                out.setdefault(cid, []).extend(parsed)

    if not out:
        raise ValueError(
            "No stakeholder rows found. Give each scope its own sheet (named as "
            "in the Scopes index), or add 'Vertical' and/or 'Region' columns to "
            "a single sheet."
        )
    return out


def _compose_config_id(vertical, region) -> str:
    v = str(vertical).strip().lower().replace(" ", "_") if vertical else ""
    r = str(region).strip().lower().replace(" ", "_") if region else ""
    if v and r:
        return f"vertical_{v}__{r}"
    if v:
        return f"vertical_{v}"
    if r:
        return r
    return "canonical"


def _rows_to_stakeholders(cols: dict, rows: list, config_id: str) -> list:
    """Turn raw rows into stakeholder dicts using the single-sheet key map."""
    out = []
    seen = set()
    for row in rows:
        def cell(name):
            i = cols.get(name)
            if i is None or i >= len(row):
                return None
            v = row[i]
            return v.strip() if isinstance(v, str) else v

        sid = cell("ID")
        if not sid:
            continue
        sid = str(sid).strip()
        if sid in seen:
            raise ValueError(f"{config_id}: duplicate stakeholder ID {sid!r}")
        seen.add(sid)

        sh = {}
        for column, key in _COLUMN_TO_KEY.items():
            if column.startswith(("Engagement Tactic", "Intel Dossier")):
                continue
            v = cell(column)
            if v is not None and v != "":
                sh[key] = v
        sh["id"] = sid

        dossier = [cell(f"Intel Dossier {i}") for i in (1, 2, 3)]
        sh["intel_dossier"] = [d for d in dossier if d]

        tactics = []
        for i in (1, 2, 3):
            raw = cell(f"Engagement Tactic {i}")
            if raw:
                t = _parse_tactic(str(raw))
                if t:
                    tactics.append(t)
        sh["engagement_tactics"] = tactics
        out.append(sh)
    return out


def main():
    if len(sys.argv) < 2:
        print("Usage: python stakeholder_config_excel.py [export|import] [path]")
        print("  export  — Generate Excel from canonical STAKEHOLDERS")
        print("  import  — Parse Excel back to stakeholder dicts (JSON to stdout)")
        sys.exit(1)

    command = sys.argv[1].lower()

    if command == "export":
        from stakeholder_map import STAKEHOLDERS
        output_path = sys.argv[2] if len(sys.argv) > 2 else "stakeholder_matrix.xlsx"
        export_stakeholders_to_excel(STAKEHOLDERS, output_path)
        print(f"[OK] Exported {len(STAKEHOLDERS)} stakeholders to: {output_path}")

    elif command == "import":
        input_path = sys.argv[2] if len(sys.argv) > 2 else "stakeholder_matrix.xlsx"
        try:
            stakeholders = import_stakeholders_from_excel(input_path)
            print(json.dumps(stakeholders, indent=2, ensure_ascii=False))
            print(f"\n[OK] Imported {len(stakeholders)} stakeholders from: {input_path}")
        except ValueError as e:
            print(f"[ERROR] {e}", file=sys.stderr)
            sys.exit(1)

    else:
        print(f"Unknown command: {command}")
        print("Use 'export' or 'import'")
        sys.exit(1)


if __name__ == "__main__":
    main()
