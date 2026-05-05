import docx
import os

def append_optimal_strategy():
    base_dir = r"c:\Users\Home\.gemini\antigravity\scratch\muressons-sim"
    source_doc = os.path.join(base_dir, "Muressons_Simulation_Briefings ver 8.docx")
    
    if not os.path.exists(source_doc):
        print(f"File not found: {source_doc}")
        return
        
    doc = docx.Document(source_doc)
    
    # Adding a Page Break to cleanly separate the new section
    doc.add_page_break()
    
    # Main Heading
    doc.add_heading('Section 15: Optimal BU Investment Strategy — Mathematical Analysis', level=1)
    
    doc.add_paragraph(
        "To provide facilitators with the mathematical 'source of truth' for the simulation's scoring engine, "
        "this section outlines the optimal capital allocation strategy across the various Business Units (BUs). "
        "The simulation is designed to computationally penalize extreme over-allocation while rewarding balanced, "
        "multi-dimensional ESG investments that maximize the final Regenerative Multiple (M_R) and Synergy indices."
    )
    
    doc.add_heading('1. The Diminishing Marginal Returns Formula', level=2)
    doc.add_paragraph(
        "Capital injected into a single BU is subject to a non-linear, diminishing returns curve. "
        "Dumping the entire Strategic Investment Pool (up to the 120% debt ceiling limit) into a single slider "
        "yields sub-optimal improvement in Social License (SLO) or Ecosystem Health Index (EHI)."
    )
    p_math1 = doc.add_paragraph()
    p_math1.add_run("Effectiveness = Base_Impact * (1 - e^(-k * Investment_Ratio))\n").font.name = 'Courier New'
    p_math1.add_run("Where k is the friction constant (typically 2.5), and Investment_Ratio is (Allocated_Capital / Ideal_Capital).").font.name = 'Courier New'
    
    doc.add_heading('2. Synergy and Cross-BU Hedging', level=2)
    doc.add_paragraph(
        "The highest-scoring strategies (yielding the 'Regenerative Titan' archetype) spread capital concurrently "
        "across multiple BUs to trigger synergy multipliers. The engine evaluates the covariance of investments "
        "across Environmental (e.g., carbon intensity reduction) and Social (e.g., burnout mitigation) axes."
    )
    p_math2 = doc.add_paragraph()
    p_math2.add_run("Synergy_Multiplier = 1.0 + (MIN(Env_Investment_%, Soc_Investment_%) * 0.5)\n").font.name = 'Courier New'
    p_math2.add_run("Example: A 50/50 split yields a 1.25x synergy multiplier on outcomes, whereas a 90/10 split yields only a 1.05x multiplier.").font.name = 'Courier New'
    
    doc.add_heading('3. The Penalty for Ignoring Systemic Debt', level=2)
    doc.add_paragraph(
        "Investments that optimize for pure revenue growth without addressing Natural Capital Debt (NCD) or "
        "transition risk actively erode the exit valuation. The backend scoring matrix applies a compound discount "
        "rate to BUs operating with a Carbon Intensity (CI) above the sectoral threshold."
    )
    p_math3 = doc.add_paragraph()
    p_math3.add_run("Exit_Valuation = (Revenue * Base_Multiple) - (NCD_Penalty + Stranded_Asset_Risk)\n").font.name = 'Courier New'
    p_math3.add_run("Therefore, the mathematically optimal move is to allocate capital to halve CI in the most pollutive BUs before expanding their revenue base.").font.name = 'Courier New'

    doc.add_heading('4. Conclusion on "Gaming" the System', level=2)
    doc.add_paragraph(
        "Because final competency scores are a 50/50 blend of these data-derived metrics and qualitative LLM-assessed "
        "interview responses (CEO Interview), the system mathematically discourages 'gaming' via single-axis maximization. "
        "An optimal financial spreadsheet without a coherent, justified stakeholder narrative will fail the final competency assessment."
    )

    output_path = os.path.join(base_dir, "Muressons_Simulation_Briefings ver 9.docx")
    doc.save(output_path)
    print(f"Successfully created {output_path}")

if __name__ == '__main__':
    append_optimal_strategy()
