import docx
import os

def append_pedagogy():
    base_dir = r"c:\Users\Home\.gemini\antigravity\scratch\muressons-sim"
    source_doc = os.path.join(base_dir, "Muressons_Simulation_Briefings ver 9.docx")
    
    if not os.path.exists(source_doc):
        print(f"File not found: {source_doc}")
        return
        
    doc = docx.Document(source_doc)
    doc.add_page_break()
    
    doc.add_heading('Section 16: Pedagogical Scaffolding & Analytics Framework', level=1)
    
    doc.add_paragraph(
        "To elevate the Muressons simulation from a mere numerical exercise into a deeply metacognitive learning experience, "
        "the Pedagogical Engine (pedagogical_engine.py) introduces several layers of cognitive scaffolding. These features "
        "are exposed through the PlayerAnalytics and PedagogicalScaffolding frontend components, heavily relying on established "
        "educational psychology frameworks."
    )
    
    doc.add_heading('1. Progressive Engine Disclosure (Fog of Complexity)', level=2)
    p1 = doc.add_paragraph()
    p1.add_run("Theory Base: Sweller (1988) Cognitive Load Theory\n").bold = True
    p1.add_run("To prevent cognitive overload, the simulation masks advanced engines (like Supply Chain Contagion or Inflation Jitter) in early rounds. "
               "The 'Fog of Complexity' lifts progressively, reflecting the increasing competency of the player.")
    
    doc.add_heading('2. Board Room Moments & Prediction Gates', level=2)
    p2 = doc.add_paragraph()
    p2.add_run("Theory Base: Moon (2004) Three-Stage Reflection & Klein (2007) Prospective Hindsight\n").bold = True
    p2.add_run("Players are periodically forced into 'Board Room Moments' asking them to structure their thinking into Noticing, Making Sense, and Working with Meaning. "
               "Prediction Gates act as pre-mortems, forcing students to state expected outcomes and later calibrating their alignment when actual results emerge.")
    
    doc.add_heading('3. Confidence Calibration & Mental Models', level=2)
    p3 = doc.add_paragraph()
    p3.add_run("Theory Base: Dunning-Kruger Effect & Argyris (1977) Double-Loop Learning\n").bold = True
    p3.add_run("At Rounds 1, 5, and 10, players rank their strategic priorities (Mental Model Tracker). The system graphs the delta in these priorities to show how their assumptions evolved. "
               "Furthermore, the Confidence Calibration engine tracks subjective certainty against objective outcome quality, actively nudging players who exhibit overconfidence.")
    
    doc.add_heading('4. Peer Comparison Nudges & Mid-Game Checkpoints', level=2)
    p4 = doc.add_paragraph()
    p4.add_run("Theory Base: Festinger (1954) Social Comparison & Vygotsky (1978) Zone of Proximal Development\n").bold = True
    p4.add_run("In multiplayer settings, blinded percentile rankings (Peer Comparison) leverage social motivation. Between rounds 5 and 6, a mid-game formative checkpoint interrupts the flow, "
               "projecting an end-state archetype based on current trajectory and prompting Vygotskian peer-to-peer advice giving.")
               
    doc.add_heading('5. Three-Phase Debrief Protocol', level=2)
    p5 = doc.add_paragraph()
    p5.add_run("Theory Base: Thiagarajan (1993)\n").bold = True
    p5.add_run("The facilitator dashboard structures the debrief into three phases: 'How Do You Feel?' (emotional offloading), 'What Happened?' (factual reconstruction using the stochastic vs strategic labeling engine), "
               "and 'So What? / Now What?' (extracting transferable real-world principles).")

    output_path = os.path.join(base_dir, "Muressons_Simulation_Briefings ver 10.docx")
    doc.save(output_path)
    print(f"Successfully created {output_path}")

if __name__ == '__main__':
    append_pedagogy()
