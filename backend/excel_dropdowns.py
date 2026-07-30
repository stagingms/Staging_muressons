"""excel_dropdowns.py — dropdown catalogues for the bulk-upload templates.

WHY THIS MODULE EXISTS
----------------------
The bulk-upload templates were plain header rows: every constrained column
(paradigm, role, ending pathway, simulation mode, industry vertical, the
boolean flags) had to be typed from memory or copied out of help text. A typo
is not caught until upload, and one bad cell fails the row — or, worse, the
value is accepted here but rejected later by a different validator.

That last failure mode is real and was found while building this: the bulk
help text advertised `un_sdg` as a paradigm, but POST /api/simulations/start
rejects it with 422. A facilitator could import a "valid" sheet and then be
unable to create a cohort.

So the catalogues below are DERIVED from the authoritative sources rather than
retyped:
  * paradigms       — config.VALID_DECISION_PARADIGMS, the same set the
                      create-session validator enforces
  * ending pathways — ending_pathways.ALL_PATHWAY_IDS
  * verticals       — materiality_db.get_bu_ids()
  * side tracks     — side_tracks.get_track_catalog()
If a new pathway or vertical ships, the dropdown gains it automatically.

Each list is resolved defensively: a template download must never 500 because
an optional catalogue failed to import.
"""
from __future__ import annotations


# ── Fixed vocabularies (no richer source exists for these) ──────────────────
ROLES = ["facilitator", "lead_facilitator"]
SIMULATION_MODES = ["conglomerate", "single_bu"]
BOOLEANS = ["TRUE", "FALSE"]


def _safe(fn, fallback):
    try:
        out = list(fn())
        return out or list(fallback)
    except Exception:
        return list(fallback)


def decision_paradigms() -> list[str]:
    def _get():
        from config import VALID_DECISION_PARADIGMS
        return sorted(VALID_DECISION_PARADIGMS)
    return _safe(_get, ["legacy_abc", "multi_toggles", "advanced_climate", "healthcare"])


def ending_pathways() -> list[str]:
    def _get():
        from ending_pathways import ALL_PATHWAY_IDS
        return list(ALL_PATHWAY_IDS)
    return _safe(_get, ["activist_ultimatum", "climate_black_swan",
                        "stakeholder_revolt", "hostile_takeover"])


def industry_verticals() -> list[str]:
    def _get():
        import materiality_db as m
        return list(m.get_bu_ids())
    return _safe(_get, ["pharma", "electronics", "consumer_goods", "software"])


def seed_bu_slots() -> list[str]:
    """The BU ids a PLAYER can be assigned to.

    Deliberately NOT industry_verticals(): a session's bu_states always carry
    the four seed slot ids (verified against a freshly created session), and
    router.VERTICAL_SLOT_MAP exists precisely to fold every vertical down to
    one of them. Offering 'chemicals' here would look plausible and match no
    business unit in the cohort.
    """
    def _get():
        from router import VERTICAL_SLOT_MAP
        return sorted(set(VERTICAL_SLOT_MAP.values()))
    return _safe(_get, ["pharma", "electronics", "consumer_goods", "software"])


def region_ids() -> list[str]:
    """Every region a cohort can be FORMED in.

    Source is regional_reporting.REGION_FRAMEWORKS — the disclosure regimes the
    engine actually models, and the same six the Create-Cohort form offers
    (frontend REGIONS in CreateCohortModal).

    This deliberately is NOT stakeholder_bulk_excel.REGION_IDS, which was the
    previous source and carries only five: it has no `china`, so a cohort formed
    in China could not have its own region named on its own roster — the exact
    dropdown-vs-server drift this module exists to prevent, pointing the other
    way. That constant stays as-is; it is an ALIAS table for parsing free-text
    stakeholder sheets ("in", "usa", "emea"), not a catalogue of what exists.
    """
    def _get():
        from regional_reporting import REGION_FRAMEWORKS
        return list(REGION_FRAMEWORKS)
    return _safe(_get, ["south_asia", "china", "europe",
                        "north_america", "asean", "africa"])


def side_track_ids() -> list[str]:
    def _get():
        from side_tracks import get_track_catalog
        cat = get_track_catalog()
        if isinstance(cat, dict):
            return list(cat.keys())
        return [t.get("track_id") for t in cat if t.get("track_id")]
    return _safe(_get, ["supply_chain", "ethics_sustainability"])


# ── Column → dropdown mapping for the facilitator template ──────────────────
# Keyed by the canonical column name so a column reorder cannot mis-target a
# validation (the builder resolves the letter from the header row).
def facilitator_column_choices() -> dict[str, list[str]]:
    return {
        "decision_paradigm": decision_paradigms(),
        "role": ROLES,
        "ending_pathway": ending_pathways(),
        "simulation_mode": SIMULATION_MODES,
        "industry_vertical": industry_verticals(),
        "shockwave_enabled": BOOLEANS,
        "trading_floor_enabled": BOOLEANS,
        "situation_room_enabled": BOOLEANS,
    }


def player_column_choices() -> dict[str, list[str]]:
    """DEPRECATED — kept for the facilitator-template call sites only.

    A player roster's columns depend on the COHORT: a 4-BU conglomerate roster
    has no per-player business unit at all, and a single-business roster names an
    industry rather than a slot. A mode-blind choice map cannot express either,
    so `roster_shape.resolve_roster_shape(session_info).choices` is the source
    for player rosters. See roster_shape.py for why the distinction is not
    cosmetic.
    """
    return {
        "assigned_bu": seed_bu_slots(),
        "region_id": region_ids(),
    }


def master_cohort_column_choices() -> dict[str, list[str]]:
    """Cohorts sheet of the master provisioning workbook."""
    return {
        "simulation_mode": SIMULATION_MODES,
        "industry_vertical": industry_verticals(),
        "decision_paradigm": decision_paradigms(),
        "region_id": region_ids(),
    }


def master_facilitator_column_choices() -> dict[str, list[str]]:
    return {"role": ROLES}


# ── Workbook helpers ────────────────────────────────────────────────────────
LIST_SHEET = "_Options"
DATA_ROWS = 500  # rows of the sheet that carry validation


def _col_letter(idx: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(idx)


def add_dropdowns(wb, ws, columns: list[str], choices: dict[str, list[str]],
                  first_row: int = 2, last_row: int = DATA_ROWS) -> None:
    """Attach list-validation dropdowns to `ws` for every column in `choices`.

    Options live on a hidden sheet and are referenced by RANGE, not inlined.
    Excel caps an inline list formula at 255 characters — the vertical and
    pathway lists are already close to that and would silently truncate (or
    corrupt the file) as the catalogues grow.
    """
    from openpyxl.worksheet.datavalidation import DataValidation

    lists = wb[LIST_SHEET] if LIST_SHEET in wb.sheetnames else wb.create_sheet(LIST_SHEET)
    lists.sheet_state = "hidden"  # present for Excel, out of the author's way

    header_to_idx = {name: i + 1 for i, name in enumerate(columns)}

    # Allocate option columns AFTER whatever is already there. The master
    # workbook calls this once per data sheet (Facilitators/Cohorts/Players);
    # restarting at column A each time would overwrite the previous sheet's
    # option lists and leave its dropdowns pointing at the wrong values.
    used = lists.max_column if lists.max_row > 1 or lists.cell(row=1, column=1).value else 0

    for offset, (field, options) in enumerate(sorted(choices.items()), start=1):
        col_no = used + offset
        if field not in header_to_idx:
            continue  # column not in this template — skip rather than mis-target
        letter = _col_letter(col_no)
        lists.cell(row=1, column=col_no, value=field)
        for r, opt in enumerate(options, start=2):
            lists.cell(row=r, column=col_no, value=opt)

        ref = f"'{LIST_SHEET}'!${letter}$2:${letter}${len(options) + 1}"
        dv = DataValidation(type="list", formula1=ref, allow_blank=True, showDropDown=False)
        # showDropDown=False is openpyxl's inverted flag: False SHOWS the arrow.
        dv.errorTitle = f"Invalid {field}"
        dv.error = ("Pick a value from the dropdown. These are the values the "
                    "simulation accepts; anything else is rejected on upload.")
        dv.promptTitle = field
        dv.prompt = "Choose from the list. Leave blank to accept the default."
        ws.add_data_validation(dv)
        tgt = _col_letter(header_to_idx[field])
        dv.add(f"{tgt}{first_row}:{tgt}{last_row}")


def add_numeric_validation(ws, columns: list[str], field: str,
                           lo: int, hi: int, first_row: int = 2,
                           last_row: int = DATA_ROWS) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation
    if field not in columns:
        return
    letter = _col_letter(columns.index(field) + 1)
    dv = DataValidation(type="whole", operator="between",
                        formula1=str(lo), formula2=str(hi), allow_blank=True)
    dv.errorTitle = f"Invalid {field}"
    dv.error = f"{field} must be a whole number between {lo} and {hi}."
    ws.add_data_validation(dv)
    dv.add(f"{letter}{first_row}:{letter}{last_row}")


def write_reference_sheet(wb, title: str = "Reference") -> None:
    """A visible cheat-sheet for the columns a dropdown cannot express —
    multi-value (side_tracks) and free-text-with-a-format (dates)."""
    from openpyxl.styles import Font
    if title in wb.sheetnames:
        return
    ref = wb.create_sheet(title)
    ref.column_dimensions["A"].width = 24
    ref.column_dimensions["B"].width = 96

    rows = [
        ("Column", "Accepted values / format"),
        ("name", "REQUIRED. Everything else is optional and defaults sensibly."),
        ("email", "Optional. Used for account notifications."),
        ("contact_number", "Optional free text."),
        ("programme", "Optional free text, e.g. 'MBA 2026'."),
        ("start_date / end_date", "YYYY-MM-DD, e.g. 2026-08-01. Blank = no window."),
        ("max_cohorts", "Whole number 1–100. Blank = 3."),
        ("decision_paradigm", ", ".join(decision_paradigms())),
        ("role", ", ".join(ROLES) + "  (a role above your own tier is ignored)"),
        ("ending_pathway", ", ".join(ending_pathways())),
        ("simulation_mode", ", ".join(SIMULATION_MODES)),
        ("industry_vertical", ", ".join(industry_verticals()) + "  (single_bu mode only)"),
        ("side_tracks", "SEMICOLON-separated ids — a dropdown cannot express multi-select. "
                        "Valid ids: " + "; ".join(side_track_ids())),
        ("shockwave_enabled", "TRUE / FALSE (blank = TRUE)"),
        ("trading_floor_enabled", "TRUE / FALSE (blank = TRUE)"),
        ("situation_room_enabled", "TRUE / FALSE (blank = TRUE)"),
    ]
    for r in rows:
        ref.append(list(r))
    for cell in ref[1]:
        cell.font = Font(bold=True)
