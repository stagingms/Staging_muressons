import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def main():
    doc = Document()

    # -- Page Setup --
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # -- Styles --
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    # Set Heading 1 Style
    h1_style = doc.styles['Heading 1']
    h1_style.font.name = 'Calibri'
    h1_style.font.size = Pt(18)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b) # Dark Slate Blue
    h1_style.paragraph_format.space_before = Pt(12)
    h1_style.paragraph_format.space_after = Pt(6)

    # Set Heading 2 Style
    h2_style = doc.styles['Heading 2']
    h2_style.font.name = 'Calibri'
    h2_style.font.size = Pt(14)
    h2_style.font.bold = True
    h2_style.font.color.rgb = RGBColor(0x47, 0x55, 0x69) # Cool Grey
    h2_style.paragraph_format.space_before = Pt(10)
    h2_style.paragraph_format.space_after = Pt(4)

    # Helper function to add tables
    def add_table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = 'Light Grid Accent 1'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = h
            for p in c.paragraphs:
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(10)
        for row_data in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row_data):
                cells[i].text = str(val)
                for p in cells[i].paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
        return t

    # ══════════════ COVER PAGE ══════════════
    for _ in range(5):
        doc.add_paragraph('')

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run('COURSE SYLLABUS')
    r_title.font.size = Pt(28)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = subtitle_p.add_run('Strategic Sustainability & Corporate Systems\n(The Muressons Simulation & BRSR Deep Dive)')
    r_sub.font.size = Pt(16)
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    r_sub.bold = True

    doc.add_paragraph('')

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = meta_p.add_run('Course Code: SUS-820  •  2nd-Year MBA Elective\n20 Sessions  •  Case-Method & Experiential Learning')
    r_meta.font.size = Pt(11)
    r_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

    for _ in range(8):
        doc.add_paragraph('')

    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = footer_p.add_run('Muressons Global Corporation Simulation Platform')
    r_foot.font.size = Pt(9.5)
    r_foot.italic = True
    r_foot.font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)

    doc.add_page_break()

    # ══════════════ 1. COURSE OVERVIEW ══════════════
    doc.add_heading('1. Course Overview', level=1)
    doc.add_paragraph(
        "Modern business leaders must navigate a landscape where corporate survival, financial performance, and environmental-social-governance (ESG) metrics are inextricably linked. This course provides 2nd-year MBA students with a hands-on, rigorous environment to practice strategic sustainability leadership."
    )
    doc.add_paragraph(
        "Using the Muressons Global Corporation Simulation as our core learning vehicle, students will take on C-suite roles (CEO, CFO, CSO, COO, CHRO) to manage a multi-billion dollar conglomerate spanning four business units (Pharma, Electronics, Consumer Goods, and Software). Over 10 simulation rounds representing 5 years of operations, student teams will make high-stakes capital allocation decisions, manage multi-stakeholder crises, and defend their strategies to a board of directors."
    )
    doc.add_paragraph(
        "This iteration of the course integrates Side Track 6: SEBI Business Responsibility and Sustainability Reporting (BRSR) NGRBC Deep Dive. Alongside the main corporate decisions, teams must comply with the 9 National Guidelines on Responsible Business Conduct (NGRBC), report on Essential vs. Leadership indicators, and secure BRSR Core assurance to unlock financial dividends and mitigate regulatory scrutiny."
    )

    # ══════════════ 2. LEARNING OBJECTIVES ══════════════
    doc.add_heading('2. Learning Objectives', level=1)
    objectives = [
        "Apply Systems Thinking: Intervene in complex systems using Meadows' Leverage Points, recognizing delay feedback loops, risk contagion, and non-linear thresholds.",
        "Conduct Double Materiality Audits: Implement ESRS 1 / CSRD guidelines, distinguishing between impact materiality and financial materiality to prioritize ESG capital budgets.",
        "Comply with SEBI BRSR Core: Navigate all 9 principles of the National Guidelines on Responsible Business Conduct, optimizing disclosure quality and managing assurance audits.",
        "Execute Scope 3 Decarbonization: Track and reduce Scope 1, 2, and 3 carbon emissions, optimizing Marginal Abatement Cost (MAC) curves and navigating green premiums.",
        "Optimize ESG Refinancing: Navigate the debt-capital markets to manage WACC using green bonds, understanding the mechanisms of the \"brown penalty\" and \"greenium.\"",
        "Manage Multi-Stakeholder Crises: Mitigate reputation damage using S-curve contagion models, balancing social license to operate (SLO) with workforce burnout and board governance.",
        "Detect and Prevent Greenwashing: Understand the financial and reputational consequences of making empty ESG claims without backing them with the required capital ratios."
    ]
    for obj in objectives:
        doc.add_paragraph(obj, style='List Bullet')

    # ══════════════ 3. ASSESSMENT & GRADING ══════════════
    doc.add_heading('3. Course Assessment & Grading', level=1)
    doc.add_paragraph(
        "All C-suite members are evaluated collectively as a unified leadership team. There are no individual role-based grade allocations, reflecting the interdependent nature of executive decision-making."
    )
    
    headers = ['Component', 'Weight', 'Description']
    rows = [
        ['Simulation Performance (Group)', '25%', 'Determined by Year 5 Terminal Equity Value, Regenerative Multiple (M_R), final corporate archetype, and BRSR disclosure score.'],
        ['Group Strategy Memos (Group)', '20%', 'Four 2-page memos submitted before rounds 3, 5, 7, and 9 outlining the strategic plan, financial forecasts, and risk mitigations.'],
        ['Boardroom Showdown (Group)', '25%', 'A live presentation in Session 20 defending the company’s 5-year performance and strategy to a panel of activist investors.'],
        ['Individual Reflection Portfolio (Individual)', '20%', 'Post-simulation synthesis analyzing key decisions, failures, systems loops, and lessons learned using Meadows\' Leverage Points.'],
        ['Class Participation & Peer Review (Individual)', '10%', 'Contribution to class debriefs and a 360-degree evaluation of C-suite team performance.']
    ]
    add_table(headers, rows)
    doc.add_paragraph('')

    # ══════════════ 4. READING MATERIALS ══════════════
    doc.add_heading('4. Core Reading Materials', level=1)
    readings = [
        "Required Simulation Documentation: Muressons Player Briefing, Muressons Student Manual (v8), and Muressons Technical Glossary with CAROIC.",
        "Meadows, D. (1999). Leverage Points: Places to Intervene in a System. Academy for Systems Change.",
        "Porter, M. E., & Kramer, M. R. (2011). Creating Shared Value. Harvard Business Review.",
        "SEBI Circular on Business Responsibility and Sustainability Reporting (BRSR) Core."
    ]
    for rd in readings:
        doc.add_paragraph(rd, style='List Bullet')

    # ══════════════ 5. SESSION OUTLINE ══════════════
    doc.add_heading('5. 20-Session Detailed Outline', level=1)

    sessions_data = [
        {
            "num": 1,
            "title": "Onboarding: Corporate Systems and Experiential Learning",
            "format": "Lecture, Team Formation, and Software Onboarding",
            "theme": "Systems Thinking and the Tragedy of the Commons",
            "concepts": [
                "Meadows' Leverage Points in a corporate context.",
                "Asymmetric ESG information and internal capital markets.",
                "Introduction to Muressons Global Corp's starting financials ($50M Treasury, WACC mechanics, and BU profiles)."
            ],
            "reading": "Muressons Player Briefing; Meadows, D. (1999) 'Leverage Points.'",
            "deliverable": "Form student C-suite teams and establish corporate charter."
        },
        {
            "num": 2,
            "title": "Simulation Round 1 (Foundations: ESG Baseline Assessment)",
            "format": "Experiential Gameplay (In-Class)",
            "theme": "Foundational risk baseline audit and capital allocation",
            "concepts": [
                "Audit depth: Surface-Level Scan vs. Deep Forensic Audit vs. Phased Rollout.",
                "CSF (Corporate Strategic Fund) matrix allocation across 5 pillars (Energy, Operations, Supply Chain, Offsetting, HR).",
                "Asymmetric baseline risk setup (avoiding future supply chain triggers)."
            ],
            "reading": "Muressons Student Manual (Chapters 1–3).",
            "deliverable": "Submit Round 1 decisions in Player Cockpit."
        },
        {
            "num": 3,
            "title": "Round 1 Debrief & Lecture: Double Materiality & ESG Governance",
            "format": "Interactive Debrief & Faculty Lecture",
            "theme": "Double Materiality (CSRD/ESRS 1 compliance)",
            "concepts": [
                "Financial Materiality (outside-in) vs. Impact Materiality (inside-out).",
                "The risk premium of governance failures under ESRS guidelines.",
                "Analyzing how Option A's decision set a hidden ticking time bomb (electronics_blindspot) in the system."
            ],
            "reading": "Muressons Student Manual (Chapter 4); EFRAG ESRS 1 Exposure Draft.",
            "deliverable": "Individual Reflection 1 (systems map of Round 1's feedback loop)."
        },
        {
            "num": 4,
            "title": "Simulation Round 2 (Double Materiality & BRSR Kickoff)",
            "format": "Experiential Gameplay (In-Class)",
            "theme": "Materiality budget allocation and initial BRSR Princples alignment",
            "concepts": [
                "CFO Materiality Gate: Full Alignment vs. Strategic Exceptions vs. CEO-Only Sign-off.",
                "BRSR Principles 1 and 7: Governance, Ethics, Transparency, and Lobbying Policy.",
                "Double Materiality Matrix minigame (target >= 90% accuracy for $2M bonus)."
            ],
            "reading": "Muressons Student Manual (Chapter 5); SEBI Principles 1 & 7 Guidelines.",
            "deliverable": "Submit Round 2 decisions and complete the Materiality Matrix minigame."
        },
        {
            "num": 5,
            "title": "Round 2 Debrief & Lecture: Scope 3 Accounting & Decarbonization",
            "format": "Interactive Debrief & Faculty Lecture",
            "theme": "GHG Protocol & Scope 3 Decarbonization",
            "concepts": [
                "Scope 1, 2, and 3 accounting boundaries and revenue-weighted carbon intensity formulas.",
                "Decarbonizing supply networks under transition cost pressure.",
                "Analyzing the ethical and compliance risks of bypassing regulatory governance."
            ],
            "reading": "Muressons Student Manual (Chapter 6); GHG Protocol Corporate Value Chain (Scope 3) Standard.",
            "deliverable": "Submit Group Strategy Memo #1 (governance & carbon baseline)."
        },
        {
            "num": 6,
            "title": "Simulation Round 3 (Scope 3 Decarbonization & BRSR Human Capital)",
            "format": "Experiential Gameplay (In-Class)",
            "theme": "Supplier switches, Green Bonds, and BRSR Human Capital standards",
            "concepts": [
                "Scope 3 pathways: Rapid Supplier Switch vs. Green Bond vs. Offset & Defer.",
                "BRSR Principles 3 and 5: Employee safety, well-being, living wages, and human rights.",
                "Green Fund Bidding Minigame: Capital bidding using Marginal Abatement Cost (MAC) curves."
            ],
            "reading": "Muressons Student Manual (Chapter 7); SEBI Principles 3 & 5 Guidelines.",
            "deliverable": "Submit Round 3 decisions and bid in the Green Fund."
        },
        {
            "num": 7,
            "title": "Round 3 Debrief & Lecture: Crisis Management & Reputation Contagion",
            "format": "Interactive Debrief & Faculty Lecture",
            "theme": "Reputation Risk Contagion and Stakeholder Fatigue",
            "concepts": [
                "S-curve sigmoid models of brand damage propagation.",
                "Stakeholder trust recovery decay formula: efficiency = 1 / (1 + 0.3 * crisis_count).",
                "Analyzing internal carbon pricing, green premiums, and Green Fund bidding outcomes."
            ],
            "reading": "Muressons Student Manual (Chapters 8 & 9); Coombs, W. T. (2014) 'Ongoing Crisis Communication.'",
            "deliverable": "Individual Assignment: Draw the Marginal Abatement Cost (MAC) curve for your portfolio BUs."
        },
        {
            "num": 8,
            "title": "Simulation Round 4 (Contagion & BRSR Environmental Stewardship)",
            "format": "Experiential Gameplay (In-Class)",
            "theme": "Labor exposé crisis, S-curve reputation contagion, and BRSR Environmental Stewardship",
            "concepts": [
                "Crisis response: Full Transparency vs. PR Containment vs. Deny & Deflect.",
                "BRSR Principles 6 and 2: Environmental stewardship, life-cycle sustainability, and circular resource efficiency.",
                "Sigmoid formula propagation of brand shocks across a conglomerate."
            ],
            "reading": "Muressons Student Manual (Chapter 10); SEBI Principle 6 Guidelines.",
            "deliverable": "Submit Round 4 decisions."
        },
        {
            "num": 9,
            "title": "Round 4 Debrief & Lecture: Physical Climate Risks & Scenario Analysis",
            "format": "Interactive Debrief & Faculty Lecture",
            "theme": "TCFD Framework and Climate Adaptation Economics",
            "concepts": [
                "Physical vs. transition climate risks under TCFD.",
                "Hard engineering vs. nature-based adaptation (resilience coefficients).",
                "WACC escalation: Central banks begin rate-tightening cycle (+200bps)."
            ],
            "reading": "Muressons Student Manual (Chapter 11); TCFD Implementation Guidelines.",
            "deliverable": "Submit Group Strategy Memo #2 (supply chain resilience and reputation recovery plan)."
        },
        {
            "num": 10,
            "title": "Simulation Round 5 (Climate: Physical Risk & BRSR Stakeholder Relations)",
            "format": "Experiential Gameplay (In-Class)",
            "theme": "Cyclone landfall resilience, Shadow Board conflicts, and BRSR Value Chain",
            "concepts": [
                "Cyclone protection: Hard Engineering vs. Nature-Based Solutions vs. Insurance Only.",
                "Shadow Board governance: Rejecting Planet, People, and Shareholder director advice.",
                "BRSR Principles 4, 8, and 9: Inclusive growth, community relations, and consumer value.",
                "SEBI greenwashing show-cause checks (unsubstantiated claims trigger $2.5M fine)."
            ],
            "reading": "Muressons Student Manual (Chapter 12); SEBI Principles 4, 8, 9 Guidelines.",
            "deliverable": "Submit Round 5 decisions."
        },
        {
            "num": 11,
            "title": "Round 5 Debrief & Lecture: Corporate Digital Responsibility & AI Ethics",
            "format": "Interactive Debrief & Faculty Lecture",
            "theme": "Algorithmic Bias and Technology Governance",
            "concepts": [
                "Corporate Digital Responsibility (CDR) framework.",
                "EU AI Act regulatory impact and compliance costs.",
                "Analyzing Shadow Board rejections and WACC changes."
            ],
            "reading": "Muressons Student Manual (Chapter 13); Lobschat, L. et al. (2020) 'What is Corporate Digital Responsibility?'",
            "deliverable": "Individual Reflection 2 (Analyzing C-suite dynamics during the Shadow Board dispute)."
        },
        {
            "num": 12,
            "title": "Simulation Round 6 (AI Bias & Integrated BRSR Core Disclosure)",
            "format": "Experiential Gameplay (In-Class)",
            "theme": "AI bias mitigation, VCM portfolio building, and final BRSR audit",
            "concepts": [
                "AI bias: Monetize Algorithm vs. Ethical AI Overhaul vs. Quiet Patch.",
                "VCM Portfolio Builder: Distributing capital across Carbon Credit Integrity Tiers.",
                "Integrated BRSR Core Report Submission & Assurance selection.",
                "Whistleblower leak checks (governance fragility triggers $2.5M audit cost)."
            ],
            "reading": "Muressons Student Manual (Chapters 14 & 15); SEBI BRSR Core Assurance framework.",
            "deliverable": "Submit Round 6 decisions and final BRSR report (Score >= 80 and Core Assurance unlocks +0.05 M_R dividend)."
        },
        {
            "num": 13,
            "title": "Round 6 Debrief & Lecture: Circular Economy & Industrial Ecology",
            "format": "Interactive Debrief & Faculty Lecture",
            "theme": "Circular Economy and the ReSOLVE Framework",
            "concepts": [
                "Transitioning from linear (take-make-waste) to circular business models.",
                "Closed-loop manufacturing and product-as-a-service models.",
                "Analyzing VCM credit integrity and group greenwashing risk."
            ],
            "reading": "Muressons Student Manual (Chapter 16); Ellen MacArthur Foundation 'Towards the Circular Economy.'",
            "deliverable": "Submit Group Strategy Memo #3 (carbon portfolio and AI ethics position paper)."
        },
        {
            "num": 14,
            "title": "Simulation Round 7 (Circularity: Circular Economy Transition)",
            "format": "Experiential Gameplay (In-Class)",
            "theme": "Circular design compliance and strategic synergy",
            "concepts": [
                "EU Waste diversion compliance: Circular Redesign vs. EPR Program vs. Waste-to-Energy.",
                "Circular Strategy Dashboard minigame (optimize resource circularity).",
                "Unlocking the strategic synergy opex reduction formula."
            ],
            "reading": "Muressons Student Manual (Chapter 17).",
            "deliverable": "Submit Round 7 decisions and complete Circular Strategy Dashboard."
        },
        {
            "num": 15,
            "title": "Round 7 Debrief & Lecture: Natural Capital & TNFD Framework",
            "format": "Interactive Debrief & Faculty Lecture",
            "theme": "Biodiversity, Water Scarcity, and Nature Risk",
            "concepts": [
                "Natural Capital Debt (NCD) and its direct surcharge on cost of debt (WACC).",
                "The TNFD LEAP assessment framework.",
                "Analyzing the math of the Synergy Engine (New_OPEX reduction formula)."
            ],
            "reading": "Muressons Student Manual (Chapter 18); TNFD Recommendations on Nature-Related Disclosures.",
            "deliverable": "Individual Assignment: Calculate the financial payback of your circularity investments."
        },
        {
            "num": 16,
            "title": "Simulation Round 8 (Blue Stress: Water Scarcity Emergency)",
            "format": "Experiential Gameplay (In-Class)",
            "theme": "Water footprints, rationing conflicts, and regulatory lobbying",
            "concepts": [
                "Water crisis: Efficient All-BU vs. Prioritize Electronics vs. Desalination Mega-Project.",
                "Policy War Room Minigame: Game-theoretic regulatory lobbying scenarios.",
                "Capital expenditure trade-offs of long-term infrastructure ($30M capex)."
            ],
            "reading": "Muressons Student Manual (Chapter 19).",
            "deliverable": "Submit Round 8 decisions and run lobbying scenarios in Policy War Room."
        },
        {
            "num": 17,
            "title": "Round 8 Debrief & Lecture: Human Capital & Just Transition",
            "format": "Interactive Debrief & Faculty Lecture",
            "theme": "Just Transition & Social License to Operate (SLO)",
            "concepts": [
                "The socio-economic impacts of decarbonization.",
                "Employee burnout index and its quadratic OPEX penalties.",
                "Workforce readiness as a multiplier for transition strategy success."
            ],
            "reading": "Muressons Student Manual (Chapter 20); ILO Guidelines for a Just Transition.",
            "deliverable": "Submit Group Strategy Memo #4 (water infrastructure and regulatory lobbying summary)."
        },
        {
            "num": 18,
            "title": "Simulation Round 9 (Just Transition: Workforce & Community)",
            "format": "Experiential Gameplay (In-Class)",
            "theme": "Factory closures, employee strike risks, and green debt auctions",
            "concepts": [
                "Closure pathways: Immediate Closure vs. Managed Transition vs. Community Fund.",
                "ESG Refinancing Simulator: Refinancing maturing debt using green bonds.",
                "Balancing the WACC brown penalty vs. greenium in bond auctions.",
                "Strike probability check (low SLO + immediate closure triggers 75% strike chance)."
            ],
            "reading": "Muressons Student Manual (Chapter 21).",
            "deliverable": "Submit Round 9 decisions and execute refinancing in the ESG Refinancing Simulator."
        },
        {
            "num": 19,
            "title": "Round 9 Debrief & Lecture: Strategic Valuation & Governance",
            "format": "Interactive Debrief & Faculty Lecture",
            "theme": "Corporate Governance, Shareholder Activism, and Valuation",
            "concepts": [
                "Dissecting the Muressons Terminal Valuation Formula: EBITDA, WACC-linked exit multiples, and M_R adjustment.",
                "Defending against hostile takeovers and activist interventions.",
                "Monitoring the Foreshadowing KPIs (Stranded Asset Exposure, Social Capital Index, Takeover Vulnerability Index, and Compliance Risk Index) to anticipate the dynamic Round 10 climax."
            ],
            "reading": "Muressons Student Manual (Chapter 22); McKinsey & Co. Valuation (Chapters on ESG).",
            "deliverable": "Submit the final Group Pitch Deck for the Boardroom Showdown."
        },
        {
            "num": 20,
            "title": "Climax: Simulation Round 10, Boardroom Showdown & Final Debrief",
            "format": "Live Capstone (3 Hours)",
            "theme": "Emergent pathway climax, Boardroom defense, and systems thinking debrief",
            "concepts": [
                "Play Round 10: Dynamic climax based on cohort metrics (Activist Ultimatum, Climate Black Swan, Stakeholder Revolt, Hostile Takeover, or Regulatory Shutdown).",
                "CEO Interview: AI-scored post-game leadership assessment across 6 dimensions.",
                "Boardroom Showdown: Live defense of C-suite performance to activist directors.",
                "Systems wrap-up: Meadows' Leverage Points applied to student outcomes."
            ],
            "reading": "Muressons Student Manual (Chapter 23).",
            "deliverable": "Live Boardroom presentation, CEO interview, and Final Reflection Portfolio."
        }
    ]

    for s in sessions_data:
        doc.add_heading(f"Session {s['num']}: {s['title']}", level=2)
        doc.add_paragraph(f"Format: {s['format']}", style='Normal').runs[0].bold = True
        doc.add_paragraph(f"Theoretical/Theme Focus: {s['theme']}", style='Normal').runs[0].italic = True
        
        doc.add_paragraph("Key Concepts Covered:")
        for c in s['concepts']:
            doc.add_paragraph(c, style='List Bullet')
            
        doc.add_paragraph(f"Required Reading: {s['reading']}", style='Normal')
        doc.add_paragraph(f"Session Deliverable: {s['deliverable']}", style='Normal')
        doc.add_paragraph('')

    # ══════════════ 6. TECHNICAL CONFIGURATIONS ══════════════
    doc.add_heading('6. Technical Configurations for Faculty', level=1)
    configs = [
        "Difficulty Tier: Set to Expert in God Mode to activate WACC interest rate escalation, regulatory sandbox constraints (SE-7), and detailed balance sheet engine calculations (IAS 1 formatting).",
        "Pacing Mode: Set to Facilitator-Paced. Open rounds sequentially to align with the syllabus sessions. Do not use Free-Play mode in a structured semester.",
        "Pillars Paradigm: Use the Strategic Pillars (multi_toggles) decision paradigm. This forces students to manage resource scarcity and balance tradeoffs across 5 independent pillars every round, in addition to the main crisis.",
        "Ending Pathway Selection: Set to Random / Dynamic (random). This forces students to monitor their foreshadowing KPIs (SAE, SCI, TVI, CRI) from Round 5 onwards, since the final exam scenario is dynamic and based on their corporate performance.",
        "Incorporating Side Tracks: Enable Side Track 6 (BRSR NGRBC Deep Dive) in God Mode and configure the Facilitator dashboard to inject ST-R1 in Round 2, ST-R2 in Round 3, ST-R3 in Round 4, ST-R4 in Round 5, and ST-R5 in Round 6."
    ]
    for cfg in configs:
        doc.add_paragraph(cfg, style='List Bullet')

    # Save
    out_dir = os.path.dirname(__file__)
    # Root of repo:
    root_dir = os.path.dirname(out_dir)
    out_file = os.path.join(root_dir, 'Muressons_MBA_Course_Syllabus.docx')
    doc.save(out_file)
    print(f"[OK] Word document successfully generated and saved to: {out_file}")

if __name__ == '__main__':
    main()
