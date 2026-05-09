from docx import Document
d = Document('Muressons_Student_Manual_v3.docx')
headings = [(i, p.text) for i, p in enumerate(d.paragraphs) if p.style.name.startswith('Heading')]
print(f"Total paragraphs: {len(d.paragraphs)}")
print(f"Total tables: {len(d.tables)}")
print(f"Total headings: {len(headings)}")
print()
# Show key section headings
for i, (idx, text) in enumerate(headings):
    clean = text.encode('ascii', 'replace').decode()
    if any(text.startswith(x) for x in ['PART','12.','13.','14.','15.','Case','App']):
        print(f"  [{idx}] {clean}")
print()
# Count glossary terms
glossary_start = None
for i, p in enumerate(d.paragraphs):
    if '14. Glossary' in p.text:
        glossary_start = i
        break
if glossary_start:
    bold_count = 0
    for p in d.paragraphs[glossary_start:glossary_start+120]:
        for r in p.runs:
            if r.bold and ':' in p.text and len(p.text) > 20:
                bold_count += 1
                break
    print(f"Glossary terms (bold entries): ~{bold_count}")
# Check case studies
case_count = sum(1 for _, t in headings if t.startswith('Case'))
print(f"Case study sections: {case_count}")
# Check references
ref_count = sum(1 for p in d.paragraphs if 'Reference:' in p.text and 'http' in p.text)
print(f"Paragraphs with reference URLs: {ref_count}")
