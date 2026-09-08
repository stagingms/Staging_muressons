"""Insert a figure into a chapter docx: picture + numbered caption + alt text.

Works both on a live chapgen.Chapter (append at the build point) and on an
existing Document (insert after a located anchor paragraph).
"""
import copy, os
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

CAP_GREY = RGBColor(0x5A, 0x6B, 0x7B)


def _set_alt(run, alt, title=None):
    """Stamp alt text onto the inline drawing so screen readers get it."""
    for docPr in run._r.iter(qn('wp:docPr')):
        docPr.set('descr', alt)
        if title:
            docPr.set('title', title)
        return True
    return False


def _clear_line_rule(p):
    """Drop any inherited w:line from the paragraph.

    The chapter body carries w:spacing w:line="276" with no w:lineRule. Word
    reads the missing rule as "auto"; LibreOffice reads it as an exact 13.8pt
    line box, which clips an inline picture down to a 0.19in strip — the whole
    figure vanishes in the PDF proof. A figure paragraph has no business
    carrying a line rule at all.
    """
    pPr = p._p.find(qn('w:pPr'))
    if pPr is None:
        return
    sp = pPr.find(qn('w:spacing'))
    if sp is None:
        return
    for attr in ('w:line', 'w:lineRule'):
        if sp.get(qn(attr)) is not None:
            del sp.attrib[qn(attr)]


def _fill_picture(p, path, alt, width_in):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    _clear_line_rule(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    _set_alt(run, alt)
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(2)
    pf.keep_with_next = True
    return p


def _fill_caption(p, text):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    _clear_line_rule(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.bold = False
    run.font.size = Pt(8.5)
    run.font.color.rgb = CAP_GREY
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(10)
    return p


def add_to_chapter(c, path, number, caption, alt, width_in=6.5):
    """chapgen.Chapter: append figure + caption at the current build point."""
    p = c._add("body")
    _fill_picture(p, path, alt, width_in)
    cp = c._add("body")
    _fill_caption(cp, f"Figure {number}  {caption}")
    return p, cp


def insert_after_para(doc, anchor, path, number, caption, alt, width_in=6.5):
    """Existing Document: insert figure + caption immediately after `anchor`."""
    shell = copy.deepcopy(anchor._p)
    for ch in list(shell):
        if ch.tag in (qn('w:r'), qn('w:hyperlink'), qn('w:bookmarkStart'), qn('w:bookmarkEnd')):
            shell.remove(ch)
    cap_el = copy.deepcopy(shell)
    anchor._p.addnext(cap_el)
    anchor._p.addnext(shell)
    pic_p = Paragraph(shell, anchor._parent)
    cap_p = Paragraph(cap_el, anchor._parent)
    _fill_picture(pic_p, path, alt, width_in)
    _fill_caption(cap_p, f"Figure {number}  {caption}")
    return pic_p, cap_p


def strip_figures(doc):
    """Remove every figure + caption previously inserted, so re-running an
    insertion script is idempotent rather than additive."""
    from docx.text.paragraph import Paragraph
    import re
    removed = 0
    for p in list(doc.element.body.iter(qn('w:p'))):
        has_pic = p.find('.//' + qn('w:drawing')) is not None
        txt = "".join(t.text or "" for t in p.iter(qn('w:t')))
        is_cap = bool(re.match(r'^Figure \d+\.\d+\s', txt))
        if has_pic or is_cap:
            parent = p.getparent()
            if parent is not None:
                parent.remove(p)
                removed += 1
    return removed
